#!/usr/bin/env python3
"""Structured-case pipeline adapter backed by the complete Xi'an modules.

File parsing and UI-shaped output are local integration concerns. Analysis
generation, framework extraction, case-library storage and GED similarity all
delegate to the unchanged source copied under ``xian_modules``.
"""
from __future__ import annotations

import hashlib
import csv
import json
import os
import re
import sys
from datetime import datetime, timezone
from pathlib import Path
from time import perf_counter
from typing import Any


WORKER_ROOT = Path(__file__).resolve().parent
XIAN_MODULE_ROOT = WORKER_ROOT / "xian_modules"


def _load_worker_environment(path: Path) -> None:
    """Load the worker-local model configuration before Xi'an modules import."""
    if not path.is_file():
        return
    for raw_line in path.read_text(encoding="utf-8-sig").splitlines():
        line = raw_line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, value = line.split("=", 1)
        key = key.removeprefix("export ").strip()
        if not re.fullmatch(r"[A-Za-z_][A-Za-z0-9_]*", key):
            continue
        value = value.strip()
        if len(value) >= 2 and value[0] == value[-1] and value[0] in {'"', "'"}:
            value = value[1:-1]
        current = os.environ.get(key)
        if current is None or not current.strip():
            os.environ[key] = value


_load_worker_environment(WORKER_ROOT / ".env")
if str(XIAN_MODULE_ROOT) not in sys.path:
    sys.path.insert(0, str(XIAN_MODULE_ROOT))

from datagraph_bank.workflow import BankCaseWorkflow  # noqa: E402
from similarity_matching import match_similar_cases  # noqa: E402
from text_generation import ensure_analysis_text  # noqa: E402


WORKER_ID = "XI_AN_STRUCTURED_CASE_PIPELINE"
WORKER_VERSION = "2.8.0-bge-m3-event-rag"
NEW_BATCH_FIELDS = ("basic_info", "customers", "transaction_features")
HISTORICAL_BATCH_FIELDS = ("basic_info", "customers", "analysis_texts")
HISTORICAL_ENRICHED_BATCH_FIELDS = (
    "basic_info", "customers", "analysis_texts", "accounts", "evidences"
)
HISTORICAL_BATCH_FIELD_ALIASES = (
    HISTORICAL_BATCH_FIELDS,
    HISTORICAL_ENRICHED_BATCH_FIELDS,
    ("basic_info", "customers", "analysis_text"),
)
MAX_BATCH_CASES = 500
ANTI_FRAUD_NEW_BATCH_FIELDS = (
    "basic_info", "customers", "accounts", "devices", "event_chain"
)
ANTI_FRAUD_HISTORICAL_BATCH_FIELDS = (
    "basic_info", "customers", "accounts", "devices", "text_analysis"
)


def _batch_header_error(headers: tuple[str, ...], recognition_mode: str) -> str:
    expected = (
        HISTORICAL_BATCH_FIELDS
        if recognition_mode == "HISTORICAL"
        else NEW_BATCH_FIELDS
    )
    if recognition_mode == "HISTORICAL" and headers == NEW_BATCH_FIELDS:
        return (
            "历史反洗钱案例必须提供已有分析文本 analysis_texts，"
            "不能使用新增案例字段 transaction_features；CSV/XLSX 表头必须严格为: "
            + ", ".join(expected)
        )
    if headers in {ANTI_FRAUD_NEW_BATCH_FIELDS, ANTI_FRAUD_HISTORICAL_BATCH_FIELDS}:
        return (
            "当前选择的是反洗钱场景，但文件表头属于反欺诈案例；"
            "请将处理场景切换为反欺诈"
        )
    return "CSV/XLSX 表头必须严格为: " + ", ".join(expected)


def _read_json(source: Path, expected_type: type, required: bool) -> Any:
    if not source.is_file():
        if required:
            raise ValueError(f"单案例处理缺少 {source.name}")
        return None
    value = json.loads(source.read_text(encoding="utf-8-sig"))
    if not isinstance(value, expected_type):
        expected = "对象" if expected_type is dict else "数组"
        raise ValueError(f"{source.name} 顶层必须是 JSON {expected}")
    return value


def load_single_case_directory(source: Path) -> dict[str, Any]:
    """Assemble the original Xi'an case contract from five role-based files."""
    if not source.is_dir():
        raise ValueError("上传案例目录不存在")
    basic_info = _read_json(source / "basic_info.json", dict, True)
    customers = _read_json(source / "customers.json", list, True)
    if not str(basic_info.get("case_id") or "").strip():
        raise ValueError("basic_info.json 必须包含非空 case_id")
    if not customers:
        raise ValueError("customers.json 至少需要一个客户")
    for index, customer in enumerate(customers, start=1):
        if not isinstance(customer, dict) or not str(customer.get("entity_id") or "").strip():
            raise ValueError(
                f"customers.json 第 {index} 个客户必须包含非空 entity_id"
            )
    analysis_texts = _read_json(source / "analysis_texts.json", dict, False)
    accounts = _read_json(source / "accounts.json", list, False)
    other_entities = _read_json(source / "other_entities.json", list, False)
    record: dict[str, Any] = {
        "basic_info": basic_info,
        "customers": customers,
        "accounts": accounts or [],
        "other_entities": other_entities or [],
    }
    if analysis_texts is not None:
        if not any(isinstance(text, str) and text.strip() for text in analysis_texts.values()):
            raise ValueError("analysis_texts.json 至少需要包含一段非空分析文本")
        record["analysis_texts"] = analysis_texts
    return record


def _decode_batch_json(value: Any, field: str, source: str, expected: type) -> Any:
    if not isinstance(value, str) or not value.strip():
        raise ValueError(f"{source} 的 {field} 不能为空")
    try:
        decoded = json.loads(value)
    except json.JSONDecodeError as exc:
        raise ValueError(f"{source} 的 {field} 不是有效 JSON") from exc
    if not isinstance(decoded, expected):
        expected_name = "对象" if expected is dict else "数组"
        raise ValueError(f"{source} 的 {field} 必须是 JSON {expected_name}")
    return decoded


def _decode_batch_analysis_texts(
    values: dict[str, Any], source: str
) -> dict[str, str]:
    field = "analysis_texts" if "analysis_texts" in values else "analysis_text"
    value = values.get(field)
    if not isinstance(value, str) or not value.strip():
        raise ValueError(f"{source} 的 {field} 不能为空")
    try:
        decoded = json.loads(value)
    except json.JSONDecodeError as exc:
        raise ValueError(f"{source} 的 {field} 不是有效 JSON") from exc
    if isinstance(decoded, str):
        decoded = {"analysis_text": decoded}
    elif isinstance(decoded, list):
        decoded = {
            f"analysis_text{index}": str(text)
            for index, text in enumerate(decoded, start=1)
            if str(text).strip()
        }
    if not isinstance(decoded, dict):
        raise ValueError(f"{source} 的 {field} 必须是 JSON 对象、字符串或数组")
    normalized = {
        str(key): text.strip()
        for key, text in decoded.items()
        if isinstance(text, str) and text.strip()
    }
    if not normalized:
        raise ValueError(f"{source} 的 {field} 至少需要一段非空报告")
    return normalized


def _batch_record(values: dict[str, Any], source: str, recognition_mode: str) -> dict[str, Any]:
    basic_info = _decode_batch_json(values.get("basic_info"), "basic_info", source, dict)
    customers = _decode_batch_json(values.get("customers"), "customers", source, list)
    case_id = str(basic_info.get("case_id") or "").strip()
    if not case_id:
        raise ValueError(f"{source} 的 basic_info.case_id 不能为空")
    if not customers:
        raise ValueError(f"{source} 的 customers 至少需要一个客户")
    for index, customer in enumerate(customers, start=1):
        if not isinstance(customer, dict) or not str(customer.get("entity_id") or "").strip():
            raise ValueError(f"{source} 的 customers 第 {index} 项必须包含 entity_id")
    record: dict[str, Any] = {
        "basic_info": basic_info,
        "customers": customers,
        "accounts": (
            _decode_batch_json(values.get("accounts"), "accounts", source, list)
            if "accounts" in values else []
        ),
        "other_entities": [],
        "evidences": (
            _decode_batch_json(values.get("evidences"), "evidences", source, list)
            if "evidences" in values else []
        ),
    }
    if recognition_mode == "HISTORICAL":
        record["analysis_texts"] = _decode_batch_analysis_texts(values, source)
    else:
        transaction_features = _decode_batch_json(
            values.get("transaction_features"), "transaction_features", source, dict
        )
        feature_case_id = str(transaction_features.get("case_id") or "").strip()
        if feature_case_id and feature_case_id != case_id:
            raise ValueError(
                f"{source} 的 transaction_features.case_id 与 basic_info.case_id 不一致"
            )
        record["transaction_features"] = transaction_features
    return record


def load_batch_case_file(source: Path, recognition_mode: str) -> list[dict[str, Any]]:
    expected_fields = (
        HISTORICAL_BATCH_FIELDS if recognition_mode == "HISTORICAL" else NEW_BATCH_FIELDS
    )
    suffix = source.suffix.lower()
    rows: list[tuple[int, dict[str, Any]]] = []
    if suffix == ".csv":
        with source.open("r", encoding="utf-8-sig", newline="") as handle:
            reader = csv.DictReader(handle)
            headers = tuple(str(name or "").strip() for name in (reader.fieldnames or []))
            valid_headers = (
                headers in HISTORICAL_BATCH_FIELD_ALIASES
                if recognition_mode == "HISTORICAL"
                else headers == expected_fields
            )
            if not valid_headers:
                raise ValueError(_batch_header_error(headers, recognition_mode))
            for row_number, row in enumerate(reader, start=2):
                if row.get(None):
                    raise ValueError(f"CSV 第 {row_number} 行字段数超过表头")
                if any(str(value or "").strip() for value in row.values()):
                    rows.append((row_number, row))
    elif suffix == ".xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise RuntimeError("XLSX 批处理依赖 openpyxl，当前 Worker 环境未安装") from exc
        workbook = load_workbook(source, read_only=True, data_only=False)
        try:
            if len(workbook.worksheets) != 1:
                raise ValueError("XLSX 必须且只能包含一个工作表")
            sheet = workbook.worksheets[0]
            iterator = sheet.iter_rows(values_only=False)
            header_cells = next(iterator, ())
            headers = tuple(str(cell.value or "").strip() for cell in header_cells)
            valid_headers = (
                headers in HISTORICAL_BATCH_FIELD_ALIASES
                if recognition_mode == "HISTORICAL"
                else headers == expected_fields
            )
            if not valid_headers:
                raise ValueError(_batch_header_error(headers, recognition_mode))
            for row_number, cells in enumerate(iterator, start=2):
                if any(cell.data_type == "f" for cell in cells):
                    raise ValueError(f"XLSX 第 {row_number} 行不允许使用公式")
                values = ["" if cell.value is None else str(cell.value) for cell in cells]
                if any(value.strip() for value in values):
                    rows.append((row_number, dict(zip(headers, values))))
        finally:
            workbook.close()
    else:
        raise ValueError("批处理只支持 .csv 或 .xlsx 文件")
    if not rows:
        raise ValueError("批处理文件中没有案例数据")
    if len(rows) > MAX_BATCH_CASES:
        raise ValueError(f"单个批处理文件最多支持 {MAX_BATCH_CASES} 个案例")
    records = [
        _batch_record(values, f"{source.name}#row={row_number}", recognition_mode)
        for row_number, values in rows
    ]
    case_ids = [str(record["basic_info"]["case_id"]).strip() for record in records]
    duplicates = sorted({case_id for case_id in case_ids if case_ids.count(case_id) > 1})
    if duplicates:
        raise ValueError("批处理文件存在重复 case_id: " + ", ".join(duplicates[:10]))
    return records


def _directory_sha256(source: Path) -> str:
    digest = hashlib.sha256()
    for name in (
        "basic_info.json",
        "customers.json",
        "analysis_texts.json",
        "accounts.json",
        "other_entities.json",
    ):
        path = source / name
        if not path.is_file():
            continue
        digest.update(name.encode("utf-8"))
        digest.update(b"\0")
        digest.update(path.read_bytes())
        digest.update(b"\0")
    return digest.hexdigest()


def _source_sha256(source: Path) -> str:
    if source.is_file():
        return hashlib.sha256(source.read_bytes()).hexdigest()
    return _directory_sha256(source)


def _merge_history_snapshot(
    postgres_paths: list[Path], prepared: list[dict[str, Any]]
) -> list[Path]:
    """Build a case-ID-deduplicated snapshot from PG and completed rows."""
    history_by_case_id: dict[str, Path] = {}
    for path in postgres_paths:
        document = json.loads(path.read_text(encoding="utf-8-sig"))
        history_case_id = str((document.get("basic_info") or {}).get("case_id") or "").strip()
        if not history_case_id:
            raise ValueError(f"PostgreSQL 历史案例缺少 basic_info.case_id: {path}")
        history_by_case_id[history_case_id] = path
    for item in prepared:
        history_case_id = str(
            (item["framework"].get("basic_info") or {}).get("case_id") or ""
        ).strip()
        if not history_case_id:
            raise ValueError("当前已完成案例缺少 basic_info.case_id")
        history_by_case_id[history_case_id] = item["finalPath"]
    return [history_by_case_id[key] for key in sorted(history_by_case_id)]


def _object(record: dict[str, Any], *names: str) -> dict[str, Any]:
    for name in names:
        value = record.get(name)
        if isinstance(value, dict):
            return dict(value)
    return {}


def _list(record: dict[str, Any], *names: str) -> list[Any]:
    for name in names:
        value = record.get(name)
        if isinstance(value, list):
            return list(value)
    return []


def _case_identity(record: dict[str, Any], index: int) -> tuple[str, str]:
    basic = _object(record, "basic_info", "basicInfo")
    case_id = str(
        basic.get("case_id")
        or basic.get("caseId")
        or record.get("case_id")
        or record.get("caseId")
        or f"CASE-UPLOAD-{index:04d}"
    ).strip()
    case_name = str(
        basic.get("case_name")
        or basic.get("caseName")
        or record.get("case_name")
        or record.get("caseName")
        or case_id
    ).strip()
    return case_id, case_name


def _workflow_input(
    record: dict[str, Any], analysis_texts: dict[str, str], index: int
) -> dict[str, Any]:
    case_id, case_name = _case_identity(record, index)
    basic_info = _object(record, "basic_info", "basicInfo")
    basic_info.setdefault("case_id", case_id)
    basic_info.setdefault("case_name", case_name)
    return {
        "basic_info": basic_info,
        "customers": _list(record, "customers", "customer_info", "customerInfo"),
        "accounts": _list(record, "accounts", "account_info", "accountInfo"),
        "other_entities": _list(record, "other_entities", "otherEntities", "entities"),
        "evidences": _list(record, "evidences", "evidence_list", "evidenceList"),
        "transaction_features": (
            _object(
                record,
                "transaction_features",
                "transactionFeatures",
                "feature_analysis",
                "featureAnalysis",
                "transactions",
            )
        ),
        "analysis_texts": analysis_texts,
    }


def _safe_name(value: str) -> str:
    normalized = re.sub(r"[^0-9A-Za-z._-]+", "_", value).strip("._")
    return normalized[:80] or "case"


def graph_snapshot(framework: dict[str, Any]) -> dict[str, Any]:
    nodes: list[dict[str, Any]] = []
    specs = (
        ("CASE", "basic_info", ("case_id",), ("case_name", "case_id")),
        ("CUSTOMER", "customers", ("entity_id",), ("customer_name", "entity_id")),
        ("ACCOUNT", "accounts", ("entity_id",), ("account_number", "holder_name", "entity_id")),
        ("ENTITY", "other_entities", ("entity_id",), ("entity_attr_1", "entity_id")),
        ("EVENT", "events", ("event_id",), ("event_name", "event_type", "event_id")),
        ("EVIDENCE", "evidences", ("evidence_id",), ("evidence_type", "evidence_id")),
    )
    for kind, field, id_fields, name_fields in specs:
        raw = framework.get(field)
        values = [raw] if field == "basic_info" else (raw or [])
        for index, value in enumerate(values, start=1):
            if not isinstance(value, dict):
                continue
            node_id = next(
                (str(value[key]) for key in id_fields if value.get(key)),
                f"{kind}-{index}",
            )
            label = next(
                (str(value[key]) for key in name_fields if value.get(key)),
                node_id,
            )
            nodes.append(
                {"id": node_id, "type": kind, "label": label, "properties": value}
            )
    edges = []
    for index, relation in enumerate(framework.get("relationships") or [], start=1):
        if not isinstance(relation, dict):
            continue
        edges.append(
            {
                "id": relation.get("relationship_id") or f"REL-{index:04d}",
                "source": relation.get("source_node_id"),
                "target": relation.get("target_node_id"),
                "type": relation.get("relationship_type") or "关联关系",
                "properties": relation,
            }
        )
    return {
        "nodes": nodes,
        "edges": edges,
        "nodeCount": len(nodes),
        "edgeCount": len(edges),
    }


def _extract_case(
    workflow: BankCaseWorkflow,
    record: dict[str, Any],
    text_result: dict[str, Any],
    index: int,
    total: int,
    run_root: Path,
    case_progress: Any = None,
) -> tuple[dict[str, Any], Path]:
    case_id, case_name = _case_identity(record, index)
    run_dir = run_root / f"{index:04d}_{_safe_name(case_id)}"

    def report_progress(stage: str, status: str, state: dict[str, Any]) -> None:
        if case_progress is None:
            return
        case_progress({
            "caseIndex": index,
            "caseTotal": total,
            "caseId": case_id,
            "caseName": case_name,
            "stage": stage,
            "status": status,
            "eventCount": len(state.get("events") or []),
            "relationshipCount": len(state.get("relationships") or []),
        })

    final_case, output_dir = workflow.run(
        _workflow_input(record, text_result["analysisTexts"], index),
        run_dir=run_dir,
        on_progress=report_progress,
    )
    return final_case.model_dump(mode="json"), output_dir / "final_case.json"


def run_pipeline(
    request: dict[str, Any],
    case_completed: Any = None,
    case_progress: Any = None,
) -> dict[str, Any]:
    pipeline_started = perf_counter()
    budget_seconds = max(
        60.0,
        min(float(request.get("pipelineBudgetSeconds") or 540), 3600.0),
    )

    def check_budget(next_stage: str) -> None:
        elapsed = perf_counter() - pipeline_started
        if elapsed >= budget_seconds:
            raise TimeoutError(
                f"案例处理已用时 {elapsed:.1f} 秒，超过 {budget_seconds:.0f} 秒预算，"
                f"未继续执行{next_stage}"
            )

    source_value = str(request.get("sourcePath") or request.get("sourceDirectory") or "").strip()
    if not source_value:
        raise ValueError("请求缺少 sourceDirectory")
    source = Path(source_value).resolve()
    mode = str(request.get("processingMode") or "SINGLE").upper()
    recognition_mode = str(request.get("recognitionMode") or "").strip().upper()
    target_stage = str(request.get("targetStage") or "SIMILARITY").strip().upper()
    if recognition_mode not in {"NEW", "HISTORICAL"}:
        raise ValueError("识别类型必须是 NEW 或 HISTORICAL")
    if target_stage not in {"UPLOAD", "REPORT", "FRAMEWORK", "SIMILARITY"}:
        raise ValueError("targetStage 必须是 UPLOAD、REPORT、FRAMEWORK 或 SIMILARITY")
    if mode == "SINGLE":
        if not source.is_dir():
            raise ValueError("单案例上传目录不存在")
        records = [load_single_case_directory(source)]
    elif mode == "BATCH":
        if not source.is_file():
            raise ValueError("批处理上传文件不存在")
        records = load_batch_case_file(source, recognition_mode)
    else:
        raise ValueError("处理方式必须是 SINGLE 或 BATCH")
    if request.get("validateOnly"):
        return {
            "status": "SUCCEEDED",
            "processingMode": mode,
            "recognitionMode": recognition_mode,
            "caseCount": len(records),
            "caseIds": [str(record["basic_info"]["case_id"]) for record in records],
            "sourceSha256": _source_sha256(source),
        }

    if target_stage == "UPLOAD":
        results = []
        for index, record in enumerate(records, start=1):
            case_id, case_name = _case_identity(record, index)
            item = {
                "caseId": case_id,
                "caseName": case_name,
                "recognitionMode": recognition_mode,
                "rawRecord": record,
                "processingRecords": [{
                    "stage": "案例上传", "status": "SUCCEEDED",
                    "message": "案例原始材料已完成校验并进入业务案例池",
                }],
            }
            results.append(item)
            if case_completed is not None:
                case_completed(item, len(results), len(records))
        return {
            "status": "SUCCEEDED", "workerId": WORKER_ID,
            "workerVersion": WORKER_VERSION, "jobId": request.get("jobId"),
            "processingMode": mode, "recognitionMode": recognition_mode,
            "targetStage": target_stage, "sourceSha256": _source_sha256(source),
            "caseCount": len(results), "results": results,
            "completedAt": datetime.now(timezone.utc).isoformat(),
        }

    job_id = str(request.get("jobId") or datetime.now().strftime("%Y%m%d_%H%M%S_%f"))
    run_root = WORKER_ROOT / "runs" / _safe_name(job_id)
    history_values = request.get("historyFiles") or []
    if str(request.get("historySource") or "POSTGRESQL").upper() != "POSTGRESQL":
        raise ValueError("历史案例库来源必须是 PostgreSQL")
    if not isinstance(history_values, list):
        raise ValueError("PostgreSQL 历史案例快照必须是文件路径数组")
    postgres_history_snapshot = []
    for value in history_values:
        path = Path(str(value)).resolve()
        if not path.is_file():
            raise ValueError(f"PostgreSQL 历史案例快照文件不存在: {path}")
        postgres_history_snapshot.append(path)
    workflow_settings = (
        request.get("frameworkSettings")
        if isinstance(request.get("frameworkSettings"), dict)
        else {}
    )
    similarity_settings = dict(
        request.get("similaritySettings")
        if isinstance(request.get("similaritySettings"), dict)
        else {}
    )
    workflow = BankCaseWorkflow(
        kb_path=XIAN_MODULE_ROOT / "data" / "kb" / "risk_event_knowledge_base.json",
        output_root=run_root / "extraction",
        llm_enabled=bool(workflow_settings.get("llmEnabled", True)),
        chat_model=workflow_settings.get("chatModel"),
        embedding_model=workflow_settings.get("embeddingModel"),
        prompt_dir=XIAN_MODULE_ROOT / "data" / "prompts",
        corenlp_url=(
            workflow_settings.get("corenlpUrl")
            or os.environ.get("CORENLP_URL")
            or "http://127.0.0.1:9002"
        ),
        corenlp_enabled=bool(workflow_settings.get("corenlpEnabled", True)),
        llm_concurrency=int(workflow_settings.get("llmConcurrency", 5)),
        service_max_attempts=int(workflow_settings.get("serviceMaxAttempts", 3)),
        corenlp_timeout=float(workflow_settings.get("corenlpTimeoutSeconds", 120)),
    )

    completed_history: list[dict[str, Any]] = []
    history_snapshot = _merge_history_snapshot(postgres_history_snapshot, [])
    results: list[dict[str, Any]] = []
    existing_count = 0
    generated_count = 0
    try:
        for index, record in enumerate(records, start=1):
            check_budget("可疑报告生成")
            report_started = perf_counter()
            text_result = ensure_analysis_text(record)
            report_seconds = perf_counter() - report_started
            existing_count += int(not text_result["generated"])
            generated_count += int(text_result["generated"])
            if target_stage == "REPORT":
                case_id, case_name = _case_identity(record, index)
                result_item = {
                    "caseId": case_id, "caseName": case_name,
                    "recognitionMode": recognition_mode, "rawRecord": record,
                    "suspiciousReport": {**text_result, "caseId": case_id},
                    "analysisReport": {**text_result, "caseId": case_id},
                }
                results.append(result_item)
                if case_completed is not None:
                    case_completed(result_item, len(results), len(records))
                continue
            check_budget("框架抽取")
            extraction_started = perf_counter()
            framework, final_path = _extract_case(
                workflow, record, text_result, index, len(records),
                run_root / "extraction", case_progress,
            )
            extraction_seconds = perf_counter() - extraction_started
            item = {
                "record": record,
                "text": text_result,
                "framework": framework,
                "finalPath": final_path,
                "snapshot": graph_snapshot(framework),
                "timings": {
                    "reportGenerationSeconds": round(report_seconds, 3),
                    "frameworkExtractionSeconds": round(extraction_seconds, 3),
                },
            }

            # Complete one case end-to-end before starting the next one. For a
            # historical upload, completed rows become available to subsequent
            # rows in this run; PostgreSQL remains the durable baseline.
            current_history = completed_history + [item]
            history_snapshot = _merge_history_snapshot(
                postgres_history_snapshot, current_history
            )
            snapshot = item["snapshot"]
            similarity = {
                "similarCases": [], "rawResult": {},
                "algorithm": "not-run", "algorithmVersion": "not-run",
                "degraded": False, "degradationReason": None,
            }
            if target_stage == "SIMILARITY":
                check_budget("相似度匹配")
                similarity_started = perf_counter()
                effective_similarity_settings = dict(similarity_settings)
                configured_timeout_value = (
                    effective_similarity_settings.get("timeoutSeconds")
                    or os.environ.get("STRUCTURED_CASE_SIMILARITY_TIMEOUT_SECONDS", 180)
                )
                try:
                    configured_similarity_timeout = float(configured_timeout_value)
                except (TypeError, ValueError):
                    configured_similarity_timeout = 180.0
                remaining_pipeline_budget = max(
                    5.0, budget_seconds - (perf_counter() - pipeline_started)
                )
                effective_similarity_settings["timeoutSeconds"] = min(
                    configured_similarity_timeout, remaining_pipeline_budget
                )
                similarity = match_similar_cases(
                    final_path, history_snapshot, effective_similarity_settings
                )
                similarity_seconds = perf_counter() - similarity_started
                item["timings"]["similarityMatchingSeconds"] = round(
                    similarity_seconds, 3
                )
            matches = similarity["similarCases"]
            basic_info = framework.get("basic_info") or {}
            case_id = str(basic_info.get("case_id") or "")
            case_name = str(basic_info.get("case_name") or case_id)
            processing_records = [
                {
                    "stage": "文本生成",
                    "status": (
                        "SUCCEEDED" if item["text"]["generated"] else "SKIPPED"
                    ),
                    "message": (
                        "未提供 analysis_texts.json，已调用原 aml-analysis-workflow 生成"
                        if item["text"]["generated"]
                        else "检测到 analysis_texts.json，按历史案例直接复用"
                    ) + f"；耗时 {item['timings']['reportGenerationSeconds']:.3f} 秒",
                },
                {
                    "stage": "框架抽取",
                    "status": "SUCCEEDED",
                    "message": (
                        "已调用原 BankCaseWorkflow，抽取 "
                        f"{snapshot['nodeCount']} 个节点、{snapshot['edgeCount']} 条关系"
                        f"；耗时 {item['timings']['frameworkExtractionSeconds']:.3f} 秒"
                    ),
                },
                *([{
                    "stage": "相似度匹配",
                    "status": "SUCCEEDED",
                    "message": (
                        (
                            "语义模型不可用，已自动降级为结构化/GED 匹配；"
                            if similarity.get("degraded")
                            else "已调用 BGE + reranker + GED 相似度匹配；"
                        )
                        + f"固定历史快照 {len(history_snapshot)} 条，返回 {len(matches)} 条"
                        f"；耗时 {item['timings']['similarityMatchingSeconds']:.3f} 秒"
                    ),
                }] if target_stage == "SIMILARITY" else []),
            ]
            result_item = {
                "caseId": case_id,
                "caseName": case_name,
                "recognitionMode": recognition_mode,
                "rawRecord": record,
                "suspiciousReport": {**item["text"], "caseId": case_id},
                "extractionResult": {
                    "generated": item["text"]["generated"],
                    "sourceType": (
                        "SYSTEM_GENERATED"
                        if item["text"]["generated"]
                        else "HISTORICAL_REUSED"
                    ),
                    "sourceLabel": (
                        "系统生成" if item["text"]["generated"] else "历史复用"
                    ),
                    "data": framework,
                },
                "analysisReport": {**item["text"], "caseId": case_id},
                "frameworkExtraction": framework,
                "graphSnapshot": snapshot,
                "caseAnalysis": {
                    "summary": (
                        f"案例包含 {len(framework.get('events') or [])} 个事件、"
                        f"{len(framework.get('relationships') or [])} 条关系，"
                        f"命中 {len(matches)} 个相似历史案例。"
                    ),
                    "similarCases": matches,
                    "similarityResult": similarity["rawResult"],
                    "similarityAlgorithm": similarity["algorithm"],
                    "similarityAlgorithmVersion": similarity["algorithmVersion"],
                    "similarityDegraded": similarity.get("degraded", False),
                    "similarityDegradationReason": similarity.get(
                        "degradationReason"
                    ),
                    "riskLevel": basic_info.get("risk_level") or "待研判",
                },
                "similarityMatch": {
                    "summary": (
                        f"命中 {len(matches)} 个相似历史案例，"
                        f"采用 {similarity['algorithm']} 算法完成匹配。"
                    ),
                    "matches": matches,
                    "algorithm": similarity["algorithm"],
                    "algorithmVersion": similarity["algorithmVersion"],
                    "degraded": similarity.get("degraded", False),
                    "degradationReason": similarity.get("degradationReason"),
                },
                "processingRecords": processing_records,
                "stageTimings": item["timings"],
                "artifacts": {"finalCase": str(item["finalPath"])},
            }
            results.append(result_item)
            if case_completed is not None:
                case_completed(result_item, len(results), len(records))
            # Every completed case becomes a candidate for the following case
            # in this batch. Persistence is streamed independently by the Java
            # callback, while this in-memory snapshot guarantees immediate
            # within-batch matching for both NEW and HISTORICAL recognition.
            completed_history.append(item)
    finally:
        workflow.close()
    total_seconds = perf_counter() - pipeline_started
    return {
        "status": "SUCCEEDED",
        "workerId": WORKER_ID,
        "workerVersion": WORKER_VERSION,
        "jobId": request.get("jobId"),
        "processingMode": mode,
        "recognitionMode": recognition_mode,
        "targetStage": target_stage,
        "sourceSha256": _source_sha256(source),
        "caseCount": len(results),
        "existingAnalysisTextCount": existing_count,
        "generatedAnalysisTextCount": generated_count,
        "historySnapshotCount": len(history_snapshot),
        "postgresHistoryCount": len(postgres_history_snapshot),
        "currentHistoricalInputCount": existing_count,
        "historySource": "POSTGRESQL",
        "performance": {
            "budgetSeconds": budget_seconds,
            "totalSeconds": round(total_seconds, 3),
            "withinBudget": total_seconds <= budget_seconds,
        },
        "moduleReuse": {
            "textGeneration": "datagraph_bank.governed_analysis.GovernedAnalysisTextGenerator",
            "frameworkExtraction": "datagraph_bank.workflow.BankCaseWorkflow.run",
            "similarityMatching": "case_similarity.match_query",
            "caseLibrary": "PostgreSQL structured_case_library snapshot",
        },
        "results": results,
        "completedAt": datetime.now(timezone.utc).isoformat(),
    }


def main() -> int:
    try:
        request = json.load(sys.stdin)
        if not isinstance(request, dict):
            raise ValueError("请求必须是 JSON 对象")
        def emit_case(case: dict[str, Any], completed: int, total: int) -> None:
            json.dump(
                {"type": "CASE_COMPLETED", "case": case, "completed": completed, "total": total},
                sys.stdout,
                ensure_ascii=False,
            )
            sys.stdout.write("\n")
            sys.stdout.flush()

        def emit_case_progress(progress: dict[str, Any]) -> None:
            json.dump(
                {"type": "CASE_STAGE", **progress},
                sys.stdout,
                ensure_ascii=False,
            )
            sys.stdout.write("\n")
            sys.stdout.flush()

        json.dump(
            run_pipeline(
                request,
                emit_case if request.get("streamResults") else None,
                emit_case_progress if request.get("streamResults") else None,
            ),
            sys.stdout,
            ensure_ascii=False,
        )
        sys.stdout.write("\n")
        return 0
    except Exception as exc:
        json.dump(
            {"status": "FAILED", "errorType": type(exc).__name__, "error": str(exc)},
            sys.stdout,
            ensure_ascii=False,
        )
        sys.stdout.write("\n")
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
