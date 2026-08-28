#!/usr/bin/env python3
"""End-to-end anti-fraud case pipeline (JSON over stdio).

Case, customer, account and device records are mapped directly. Only risk
events and graph relationships are delegated to the LLM extraction framework.
New cases first pass through the channel rule library to form a risk-event
chain, then use the retained aml-analysis-workflow with fraud prompts.
"""
from __future__ import annotations

import csv
import hashlib
import json
import os
import re
import sys
import tempfile
from datetime import datetime, timezone
from pathlib import Path
from time import perf_counter
from typing import Any


WORKER_ROOT = Path(__file__).resolve().parent
SHARED_WORKER = WORKER_ROOT.parent / "structured-case-identification"
XIAN_MODULE_ROOT = SHARED_WORKER / "xian_modules"


def _load_environment(path: Path) -> None:
    if not path.is_file():
        return
    for raw in path.read_text(encoding="utf-8-sig").splitlines():
        line = raw.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, value = line.split("=", 1)
        key = key.removeprefix("export ").strip()
        if re.fullmatch(r"[A-Za-z_][A-Za-z0-9_]*", key) and not os.environ.get(key):
            os.environ[key] = value.strip().strip("\"'")


_load_environment(WORKER_ROOT / ".env")
_load_environment(SHARED_WORKER / ".env")
for module_root in (WORKER_ROOT, SHARED_WORKER, XIAN_MODULE_ROOT):
    if str(module_root) not in sys.path:
        sys.path.insert(0, str(module_root))

from channel import CHANNELS, infer_channel  # noqa: E402
from datagraph_bank.workflow import BankCaseWorkflow  # noqa: E402
from report_generation import generate_analysis  # noqa: E402
from risk_chain import FraudRuleEngine  # noqa: E402
from similarity_matching import match_similar_cases  # noqa: E402


WORKER_ID = "ANTI_FRAUD_CASE_PIPELINE"
WORKER_VERSION = "1.1.0"
WORKFLOW_ID = "ANTI_FRAUD_CASE_PIPELINE"
RULE_PATH = WORKER_ROOT / "data" / "kb" / "fraud_rules.json"
RISK_KB_PATH = WORKER_ROOT / "data" / "kb" / "fraud_risk_event_knowledge_base.json"
RISK_KB_COLLECTION = os.environ.get(
    "ANTI_FRAUD_QDRANT_COLLECTION", "fraud_risk_event_knowledge_base"
)

RELATIONSHIP_TYPE_CONTRACT = {
    "structural": ("涉及关系", "包含关系", "来源关系"),
    "entity_entity": ("持有关系", "社会关系"),
    "entity_event": ("参与关系", "涉及关系"),
    "event_event": ("顺承关系", "上下位关系", "应对关系"),
}
RELATIONSHIP_LAYER_ORDER = (
    "structural", "entity_entity", "entity_event", "event_event"
)


def _risk_kb_summary() -> dict[str, Any]:
    value = json.loads(RISK_KB_PATH.read_text(encoding="utf-8-sig"))
    metadata = value.get("metadata") if isinstance(value.get("metadata"), dict) else {}
    event_types = value.get("event_types") if isinstance(value.get("event_types"), list) else []
    return {
        "name": str(metadata.get("name") or ""),
        "version": str(metadata.get("version") or ""),
        "event_count": len(event_types),
        "path": str(RISK_KB_PATH),
        "qdrant_collection": RISK_KB_COLLECTION,
    }


def _read_json(path: Path, expected: type, required: bool = True) -> Any:
    if not path.is_file():
        if required:
            raise ValueError(f"单案例处理缺少 {path.name}")
        return None
    value = json.loads(path.read_text(encoding="utf-8-sig"))
    if not isinstance(value, expected):
        expected_name = "对象" if expected is dict else "数组"
        raise ValueError(f"{path.name} 顶层必须是 JSON {expected_name}")
    return value


def _normalize_devices(devices: list[dict[str, Any]]) -> list[dict[str, Any]]:
    normalized = []
    seen: set[str] = set()
    for index, device in enumerate(devices, start=1):
        if not isinstance(device, dict):
            raise ValueError(f"devices.json 第 {index} 项必须是对象")
        device_id = str(
            device.get("设备号") or device.get("device_id") or device.get("entity_id") or ""
        ).strip()
        if not device_id:
            raise ValueError(f"devices.json 第 {index} 项缺少设备号")
        if device_id in seen:
            raise ValueError(f"devices.json 存在重复设备号: {device_id}")
        seen.add(device_id)
        normalized.append(
            {
                "entity_id": device_id,
                "entity_attr_1": device_id,
                "entity_attr_2": str(device.get("设备名称") or device.get("device_name") or ""),
                "entity_attr_3": str(device.get("ip地址") or device.get("ip_address") or ""),
                "entity_attr_4": "设备",
            }
        )
    return normalized


def load_case_directory(source: Path, recognition_mode: str) -> dict[str, Any]:
    if not source.is_dir():
        raise ValueError("反欺诈案例上传目录不存在")
    basic = _read_json(source / "basic_info.json", dict)
    customers = _read_json(source / "customers.json", list)
    accounts = _read_json(source / "accounts.json", list)
    devices = _read_json(source / "devices.json", list)
    if not str(basic.get("case_id") or "").strip():
        raise ValueError("basic_info.json 必须包含非空 case_id")
    if not customers:
        raise ValueError("customers.json 至少需要一个客户")
    for index, customer in enumerate(customers, start=1):
        if not isinstance(customer, dict) or not str(customer.get("entity_id") or "").strip():
            raise ValueError(f"customers.json 第 {index} 项必须包含非空 entity_id")
    for index, account in enumerate(accounts, start=1):
        if not isinstance(account, dict) or not str(account.get("entity_id") or "").strip():
            raise ValueError(f"accounts.json 第 {index} 项必须包含非空 entity_id")
    _normalize_devices(devices)

    event_chain = _read_json(source / "event_chain.json", list, required=False)
    text_analysis = _read_json(source / "text_analysis.json", dict, required=False)
    if recognition_mode == "HISTORICAL" and text_analysis is None:
        raise ValueError("历史反欺诈案例必须包含 text_analysis.json")
    if text_analysis is not None and not str(
        text_analysis.get("text") or text_analysis.get("analysis_text") or ""
    ).strip():
        raise ValueError("text_analysis.json 必须包含非空 text")

    channel = infer_channel(basic, event_chain or [])
    basic = dict(basic)
    # Keep the user-facing Chinese field requested by the contract and an
    # English alias for downstream systems that cannot address Chinese keys.
    basic["渠道"] = channel
    basic["channel"] = channel
    record = {
        "basic_info": basic,
        "customers": customers,
        "accounts": accounts,
        "devices": devices,
        "other_entities": _normalize_devices(devices),
        "event_chain": event_chain or [],
    }
    if text_analysis is not None:
        record["text_analysis"] = text_analysis
    return record


def _decode_batch_cell(value: Any, field: str, row_number: int) -> Any:
    if value is None or (isinstance(value, str) and not value.strip()):
        raise ValueError(f"第 {row_number} 行 {field} 不能为空")
    if not isinstance(value, str):
        return value
    try:
        return json.loads(value)
    except json.JSONDecodeError as exc:
        raise ValueError(f"第 {row_number} 行 {field} 不是有效 JSON: {exc.msg}") from exc


def _batch_record(values: dict[str, Any], row_number: int, recognition_mode: str) -> dict[str, Any]:
    basic = _decode_batch_cell(values.get("basic_info"), "basic_info", row_number)
    customers = _decode_batch_cell(values.get("customers"), "customers", row_number)
    accounts = _decode_batch_cell(values.get("accounts"), "accounts", row_number)
    devices = _decode_batch_cell(values.get("devices"), "devices", row_number)
    event_chain = (
        _decode_batch_cell(values.get("event_chain"), "event_chain", row_number)
        if recognition_mode == "NEW" else []
    )
    text_analysis = (
        _decode_batch_cell(values.get("text_analysis"), "text_analysis", row_number)
        if recognition_mode == "HISTORICAL" else None
    )
    if not isinstance(basic, dict) or not str(basic.get("case_id") or "").strip():
        raise ValueError(f"第 {row_number} 行 basic_info 必须是包含非空 case_id 的 JSON 对象")
    if not isinstance(customers, list) or not customers:
        raise ValueError(f"第 {row_number} 行 customers 必须是非空 JSON 数组")
    if any(not isinstance(item, dict) or not str(item.get("entity_id") or "").strip() for item in customers):
        raise ValueError(f"第 {row_number} 行 customers 每项必须包含非空 entity_id")
    if not isinstance(accounts, list):
        raise ValueError(f"第 {row_number} 行 accounts 必须是 JSON 数组")
    if any(not isinstance(item, dict) or not str(item.get("entity_id") or "").strip() for item in accounts):
        raise ValueError(f"第 {row_number} 行 accounts 每项必须包含非空 entity_id")
    if not isinstance(devices, list):
        raise ValueError(f"第 {row_number} 行 devices 必须是 JSON 数组")
    _normalize_devices(devices)
    if recognition_mode == "NEW" and (not isinstance(event_chain, list) or not event_chain):
        raise ValueError(f"第 {row_number} 行 event_chain 必须是非空 JSON 数组")
    if recognition_mode == "HISTORICAL" and (
        not isinstance(text_analysis, dict)
        or not str(text_analysis.get("text") or text_analysis.get("analysis_text") or "").strip()
    ):
        raise ValueError(f"第 {row_number} 行 text_analysis 必须是包含非空 text 或 analysis_text 的 JSON 对象")
    channel = infer_channel(basic, event_chain)
    normalized_basic = dict(basic)
    normalized_basic["渠道"] = channel
    normalized_basic["channel"] = channel
    record = {
        "basic_info": normalized_basic,
        "customers": customers,
        "accounts": accounts,
        "devices": devices,
        "other_entities": _normalize_devices(devices),
        "event_chain": event_chain,
    }
    if text_analysis is not None:
        record["text_analysis"] = text_analysis
    return record


def load_batch_case_file(source: Path, recognition_mode: str) -> list[dict[str, Any]]:
    expected = (
        ("basic_info", "customers", "accounts", "devices", "event_chain")
        if recognition_mode == "NEW"
        else ("basic_info", "customers", "accounts", "devices", "text_analysis")
    )
    rows: list[dict[str, Any]] = []
    suffix = source.suffix.lower()
    if suffix == ".csv":
        with source.open("r", encoding="utf-8-sig", newline="") as handle:
            reader = csv.DictReader(handle)
            headers = tuple(reader.fieldnames or ())
            if headers != expected:
                raise ValueError("CSV 表头必须严格为: " + ", ".join(expected))
            rows = [dict(row) for row in reader]
    elif suffix == ".xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise ValueError("XLSX 批处理需要安装 openpyxl") from exc
        workbook = load_workbook(source, read_only=True, data_only=True)
        try:
            sheet = workbook.active
            values = sheet.iter_rows(values_only=True)
            headers = tuple(str(value or "").strip() for value in next(values, ()))
            if headers != expected:
                raise ValueError("XLSX 表头必须严格为: " + ", ".join(expected))
            rows = [dict(zip(headers, row)) for row in values if any(value not in (None, "") for value in row)]
        finally:
            workbook.close()
    else:
        raise ValueError("反欺诈批处理仅支持 CSV 或 XLSX 文件")
    if not rows:
        raise ValueError("批处理文件至少需要一条案例数据")
    records = [_batch_record(row, index, recognition_mode) for index, row in enumerate(rows, start=2)]
    case_ids = [str(record["basic_info"]["case_id"]).strip() for record in records]
    duplicates = sorted({case_id for case_id in case_ids if case_ids.count(case_id) > 1})
    if duplicates:
        raise ValueError("批处理存在重复 case_id: " + ", ".join(duplicates))
    return records


def _source_sha256(source: Path) -> str:
    digest = hashlib.sha256()
    if source.is_file():
        digest.update(source.read_bytes())
        return digest.hexdigest()
    for path in sorted(source.glob("*.json"), key=lambda item: item.name):
        digest.update(path.name.encode("utf-8"))
        digest.update(b"\0")
        digest.update(path.read_bytes())
        digest.update(b"\0")
    return digest.hexdigest()


def _safe_name(value: str) -> str:
    return re.sub(r"[^0-9A-Za-z._-]+", "_", value).strip("._")[:80] or "case"


def _workflow_input(record: dict[str, Any], analysis: dict[str, Any]) -> dict[str, Any]:
    return {
        "basic_info": record["basic_info"],
        "customers": record["customers"],
        "accounts": record["accounts"],
        "other_entities": record["other_entities"],
        "evidences": [],
        "transaction_features": {},
        "analysis_texts": analysis["analysisTexts"],
    }


def _read_extraction_audit(output_dir: Path) -> dict[str, Any]:
    audit: dict[str, Any] = {}
    for file_name, keys in (
        ("06_event_extraction.json", ("event_extraction_mode", "event_extraction_calls")),
        ("07_relationship_extraction.json", (
            "relationship_extraction_order", "relationship_extraction_calls"
        )),
    ):
        path = output_dir / file_name
        if not path.is_file():
            continue
        value = json.loads(path.read_text(encoding="utf-8-sig"))
        if isinstance(value, dict):
            audit.update({key: value.get(key) for key in keys})
    return audit


def _restore_direct_mappings(
    framework: dict[str, Any], record: dict[str, Any], audit: dict[str, Any]
) -> dict[str, Any]:
    framework["basic_info"] = record["basic_info"]
    framework["customers"] = record["customers"]
    framework["accounts"] = record["accounts"]
    framework["devices"] = record["devices"]
    framework["other_entities"] = record["other_entities"]
    metadata = framework.setdefault("processing_metadata", {})
    metadata["direct_mapped_fields"] = [
        "basic_info",
        "customers",
        "accounts",
        "devices",
    ]
    metadata["llm_extracted_fields"] = ["events", "relationships"]
    metadata["risk_event_knowledge_base_audit"] = _risk_kb_summary()
    metadata["relationship_type_contract"] = {
        stage: list(types) for stage, types in RELATIONSHIP_TYPE_CONTRACT.items()
    }
    metadata["relationship_layer_order"] = list(RELATIONSHIP_LAYER_ORDER)
    metadata.update({key: value for key, value in audit.items() if value is not None})

    actual_order = [
        str(item.get("stage") or "")
        for item in metadata.get("relationship_extraction_order") or []
        if isinstance(item, dict)
    ]
    if actual_order and actual_order != list(RELATIONSHIP_LAYER_ORDER):
        raise ValueError(f"关系分层抽取顺序不符合契约: {actual_order}")
    allowed_types = {
        relation_type
        for values in RELATIONSHIP_TYPE_CONTRACT.values()
        for relation_type in values
    }
    unexpected = sorted({
        str(item.get("relationship_type") or "")
        for item in framework.get("relationships") or []
        if isinstance(item, dict) and str(item.get("relationship_type") or "") not in allowed_types
    })
    if unexpected:
        raise ValueError("反欺诈关系类型不符合共享关系契约: " + ", ".join(unexpected))
    return framework


def graph_snapshot(framework: dict[str, Any]) -> dict[str, Any]:
    nodes: list[dict[str, Any]] = []
    specs = (
        ("CASE", [framework.get("basic_info") or {}], "case_id", "case_name"),
        ("CUSTOMER", framework.get("customers") or [], "entity_id", "customer_name"),
        ("ACCOUNT", framework.get("accounts") or [], "entity_id", "account_number"),
        ("EVENT", framework.get("events") or [], "event_id", "event_name"),
        ("EVIDENCE", framework.get("evidences") or [], "evidence_id", "evidence_type"),
    )
    for kind, values, id_key, label_key in specs:
        for index, value in enumerate(values, start=1):
            if not isinstance(value, dict):
                continue
            node_id = str(value.get(id_key) or f"{kind}-{index}")
            nodes.append(
                {
                    "id": node_id,
                    "type": kind,
                    "label": str(value.get(label_key) or node_id),
                    "properties": value,
                }
            )
    for index, device in enumerate(framework.get("devices") or [], start=1):
        device_id = str(device.get("设备号") or device.get("device_id") or f"DEVICE-{index}")
        nodes.append(
            {
                "id": device_id,
                "type": "DEVICE",
                "label": str(device.get("设备名称") or device_id),
                "properties": device,
            }
        )
    edges = [
        {
            "id": relation.get("relationship_id") or f"REL-{index:04d}",
            "source": relation.get("source_node_id"),
            "target": relation.get("target_node_id"),
            "type": relation.get("relationship_type") or "关联关系",
            "properties": relation,
        }
        for index, relation in enumerate(framework.get("relationships") or [], start=1)
        if isinstance(relation, dict)
    ]
    return {"nodes": nodes, "edges": edges, "nodeCount": len(nodes), "edgeCount": len(edges)}


def _history_paths(request: dict[str, Any]) -> list[Path]:
    values = request.get("historyFiles") or []
    if not isinstance(values, list):
        raise ValueError("历史案例快照必须是路径数组")
    paths = [Path(str(value)).resolve() for value in values]
    missing = [str(path) for path in paths if not path.is_file()]
    if missing:
        raise ValueError(f"历史案例快照不存在: {missing[0]}")
    return paths


def run_pipeline(request: dict[str, Any], case_completed: Any = None) -> dict[str, Any]:
    started = perf_counter()
    source = Path(str(request.get("sourcePath") or "")).resolve()
    mode = str(request.get("processingMode") or "SINGLE").upper()
    recognition_mode = str(request.get("recognitionMode") or "").upper()
    target_stage = str(request.get("targetStage") or "SIMILARITY").upper()
    if recognition_mode not in {"NEW", "HISTORICAL"}:
        raise ValueError("识别类型必须是 NEW 或 HISTORICAL")
    if target_stage not in {"UPLOAD", "REPORT", "FRAMEWORK", "SIMILARITY"}:
        raise ValueError("targetStage 必须是 UPLOAD、REPORT、FRAMEWORK 或 SIMILARITY")
    if mode == "BATCH":
        records = load_batch_case_file(source, recognition_mode)
        source_sha = _source_sha256(source)
        if request.get("validateOnly"):
            kb_summary = _risk_kb_summary()
            return {
                "status": "SUCCEEDED", "workflow": WORKFLOW_ID,
                "processingMode": mode, "recognitionMode": recognition_mode,
                "caseCount": len(records),
                "caseIds": [str(record["basic_info"]["case_id"]) for record in records],
                "sourceSha256": source_sha,
                "riskKnowledgeBaseVersion": kb_summary["version"],
                "riskKnowledgeBaseEventCount": kb_summary["event_count"],
            }
        batch_results: list[dict[str, Any]] = []
        temp_parent = WORKER_ROOT / "runs"
        temp_parent.mkdir(parents=True, exist_ok=True)
        for index, record in enumerate(records, start=1):
            with tempfile.TemporaryDirectory(prefix="fraud-batch-", dir=temp_parent) as directory:
                case_root = Path(directory)
                for name, key in (
                    ("basic_info.json", "basic_info"), ("customers.json", "customers"),
                    ("accounts.json", "accounts"), ("devices.json", "devices"),
                ):
                    (case_root / name).write_text(
                        json.dumps(record[key], ensure_ascii=False), encoding="utf-8")
                if recognition_mode == "NEW":
                    (case_root / "event_chain.json").write_text(
                        json.dumps(record["event_chain"], ensure_ascii=False), encoding="utf-8")
                else:
                    (case_root / "text_analysis.json").write_text(
                        json.dumps(record["text_analysis"], ensure_ascii=False), encoding="utf-8")
                child_request = dict(request)
                child_request.update({
                    "sourcePath": str(case_root), "processingMode": "SINGLE",
                    "jobId": f"{request.get('jobId') or 'BATCH'}-{index}", "validateOnly": False,
                })
                child = run_pipeline(child_request)
                result = child["results"][0]
                batch_results.append(result)
                if case_completed is not None:
                    case_completed(result, len(batch_results), len(records))
        return {
            "status": "SUCCEEDED", "workerId": WORKER_ID, "workerVersion": WORKER_VERSION,
            "workflow": WORKFLOW_ID, "jobId": request.get("jobId"),
            "processingMode": mode, "recognitionMode": recognition_mode,
            "targetStage": target_stage, "caseCount": len(batch_results),
            "sourceSha256": source_sha, "results": batch_results,
            "performance": {"totalSeconds": round(perf_counter() - started, 3)},
            "completedAt": datetime.now(timezone.utc).isoformat(),
        }
    if mode != "SINGLE":
        raise ValueError("processingMode 必须是 SINGLE 或 BATCH")
    record = load_case_directory(source, recognition_mode)
    rule_engine = FraudRuleEngine(RULE_PATH)
    risk_chain = (
        rule_engine.build(record["basic_info"]["渠道"], record["event_chain"])
        if recognition_mode == "NEW"
        else []
    )
    record["risk_event_chain"] = risk_chain
    if request.get("validateOnly"):
        kb_summary = _risk_kb_summary()
        return {
            "status": "SUCCEEDED",
            "workflow": WORKFLOW_ID,
            "processingMode": mode,
            "recognitionMode": recognition_mode,
            "caseCount": 1,
            "caseIds": [record["basic_info"]["case_id"]],
            "channel": record["basic_info"]["渠道"],
            "allowedChannels": list(CHANNELS),
            "riskEventCount": len(risk_chain),
            "ruleLibraryVersion": rule_engine.version,
            "riskKnowledgeBaseVersion": kb_summary["version"],
            "riskKnowledgeBaseEventCount": kb_summary["event_count"],
            "sourceSha256": _source_sha256(source),
        }

    case_id = str(record["basic_info"]["case_id"])
    case_name = str(record["basic_info"].get("case_name") or case_id)
    if target_stage == "UPLOAD":
        result = {
            "caseId": case_id, "caseName": case_name,
            "recognitionMode": recognition_mode, "rawRecord": record,
            "channel": record["basic_info"]["渠道"],
            "riskEventChain": risk_chain, "ruleLibraryVersion": rule_engine.version,
            "processingRecords": [{"stage": "案例上传", "status": "SUCCEEDED",
                "message": "案例原始材料已完成校验并进入业务案例池"}],
        }
        if case_completed is not None:
            case_completed(result, 1, 1)
        return {"status": "SUCCEEDED", "workerId": WORKER_ID,
            "workerVersion": WORKER_VERSION, "workflow": WORKFLOW_ID,
            "jobId": request.get("jobId"), "processingMode": mode,
            "recognitionMode": recognition_mode, "targetStage": target_stage,
            "caseCount": 1, "sourceSha256": _source_sha256(source),
            "results": [result], "completedAt": datetime.now(timezone.utc).isoformat()}

    job_id = str(request.get("jobId") or datetime.now().strftime("%Y%m%d_%H%M%S_%f"))
    run_root = WORKER_ROOT / "runs" / _safe_name(job_id)
    run_root.mkdir(parents=True, exist_ok=True)
    analysis_started = perf_counter()
    analysis = generate_analysis(record)
    analysis_seconds = perf_counter() - analysis_started
    if target_stage == "REPORT":
        result = {
            "caseId": case_id, "caseName": case_name,
            "recognitionMode": recognition_mode, "rawRecord": record,
            "channel": record["basic_info"]["渠道"],
            "riskEventChain": risk_chain, "ruleLibraryVersion": rule_engine.version,
            "suspiciousReport": {**analysis, "caseId": case_id},
            "analysisReport": {**analysis, "caseId": case_id},
        }
        if case_completed is not None:
            case_completed(result, 1, 1)
        return {"status": "SUCCEEDED", "workerId": WORKER_ID,
            "workerVersion": WORKER_VERSION, "workflow": WORKFLOW_ID,
            "jobId": request.get("jobId"), "processingMode": mode,
            "recognitionMode": recognition_mode, "targetStage": target_stage,
            "caseCount": 1, "sourceSha256": _source_sha256(source),
            "results": [result], "completedAt": datetime.now(timezone.utc).isoformat()}

    settings = request.get("frameworkSettings") if isinstance(request.get("frameworkSettings"), dict) else {}
    workflow = BankCaseWorkflow(
        kb_path=RISK_KB_PATH,
        kb_collection=RISK_KB_COLLECTION,
        output_root=run_root / "extraction",
        llm_enabled=bool(settings.get("llmEnabled", True)),
        corenlp_enabled=bool(settings.get("corenlpEnabled", True)),
        corenlp_url=str(settings.get("corenlpUrl") or os.environ.get("CORENLP_URL") or "http://127.0.0.1:9002"),
        corenlp_timeout=float(settings.get("corenlpTimeoutSeconds", 120)),
        llm_concurrency=int(settings.get("llmConcurrency", 5)),
        service_max_attempts=int(settings.get("serviceMaxAttempts", 3)),
        prompt_dir=XIAN_MODULE_ROOT / "data" / "prompts",
    )
    try:
        extraction_started = perf_counter()
        final_case, output_dir = workflow.run(
            _workflow_input(record, analysis),
            run_dir=run_root / "extraction" / _safe_name(record["basic_info"]["case_id"]),
        )
        framework = _restore_direct_mappings(
            final_case.model_dump(mode="json"),
            record,
            _read_extraction_audit(output_dir),
        )
        final_path = output_dir / "final_case.json"
        final_path.write_text(json.dumps(framework, ensure_ascii=False, indent=2), encoding="utf-8")
        extraction_seconds = perf_counter() - extraction_started
    finally:
        workflow.close()

    similarity = {"similarCases": [], "rawResult": {}, "algorithm": "not-run",
                  "algorithmVersion": "not-run", "degraded": False,
                  "degradationReason": None}
    similarity_seconds = 0.0
    if target_stage == "SIMILARITY":
        similarity_started = perf_counter()
        similarity = match_similar_cases(
            final_path,
            _history_paths(request),
            request.get("similaritySettings") if isinstance(request.get("similaritySettings"), dict) else {},
        )
        similarity_seconds = perf_counter() - similarity_started
    snapshot = graph_snapshot(framework)
    matches = similarity["similarCases"][:5]
    result = {
        "caseId": case_id,
        "caseName": str(record["basic_info"].get("case_name") or case_id),
        "recognitionMode": recognition_mode,
        "rawRecord": record,
        "channel": record["basic_info"]["渠道"],
        "riskEventChain": risk_chain,
        "ruleLibraryVersion": rule_engine.version,
        "riskKnowledgeBase": _risk_kb_summary(),
        "suspiciousReport": {**analysis, "caseId": case_id},
        "analysisReport": {**analysis, "caseId": case_id},
        "extractionResult": {
            "generated": analysis["generated"],
            "sourceType": "RISK_EVENT_CHAIN_GENERATED" if analysis["generated"] else "HISTORICAL_REUSED",
            "sourceLabel": "风险事件链生成" if analysis["generated"] else "历史复用",
            "data": framework,
        },
        "frameworkExtraction": framework,
        "graphSnapshot": snapshot,
        "caseAnalysis": {
            "summary": f"案例包含 {len(framework.get('events') or [])} 个事件、{len(framework.get('relationships') or [])} 条关系，返回 {len(matches)} 个相似案例。",
            "similarCases": matches,
            "similarityResult": similarity["rawResult"],
            "similarityAlgorithm": similarity["algorithm"],
            "similarityAlgorithmVersion": similarity["algorithmVersion"],
            "riskLevel": record["basic_info"].get("risk_level") or "待研判",
        },
        "similarityMatch": {
            "matches": matches,
            "algorithm": similarity["algorithm"],
            "algorithmVersion": similarity["algorithmVersion"],
            "degraded": similarity.get("degraded", False),
            "degradationReason": similarity.get("degradationReason"),
        },
        "processingRecords": [
            {"stage": "风险事件链", "status": "SUCCEEDED" if recognition_mode == "NEW" else "SKIPPED", "message": f"渠道={record['basic_info']['渠道']}，规则库 {rule_engine.version}，生成 {len(risk_chain)} 个风险事件"},
            {"stage": "分析文本", "status": "SUCCEEDED", "message": f"{'aml-analysis-workflow 反欺诈提示词生成' if analysis['generated'] else '历史 text_analysis.json 直接复用'}；耗时 {analysis_seconds:.3f} 秒"},
            {"stage": "框架抽取", "status": "SUCCEEDED", "message": f"基础四类信息直接映射，LLM 抽取事件和关系；耗时 {extraction_seconds:.3f} 秒"},
            *([{"stage": "相似度匹配", "status": "SUCCEEDED", "message": f"复用现有 BGE/Reranker/GED 算法，返回 Top {len(matches)}；耗时 {similarity_seconds:.3f} 秒"}] if target_stage == "SIMILARITY" else []),
        ],
        "artifacts": {"finalCase": str(final_path)},
    }
    if case_completed is not None:
        case_completed(result, 1, 1)
    return {
        "status": "SUCCEEDED",
        "workerId": WORKER_ID,
        "workerVersion": WORKER_VERSION,
        "workflow": WORKFLOW_ID,
        "jobId": request.get("jobId"),
        "processingMode": mode,
        "recognitionMode": recognition_mode,
        "targetStage": target_stage,
        "caseCount": 1,
        "sourceSha256": _source_sha256(source),
        "performance": {"totalSeconds": round(perf_counter() - started, 3)},
        "results": [result],
        "completedAt": datetime.now(timezone.utc).isoformat(),
    }


def main() -> int:
    try:
        request = json.load(sys.stdin)
        if not isinstance(request, dict):
            raise ValueError("请求必须是 JSON 对象")

        def emit(case: dict[str, Any], completed: int, total: int) -> None:
            json.dump(
                {"type": "CASE_COMPLETED", "case": case, "completed": completed, "total": total},
                sys.stdout,
                ensure_ascii=False,
            )
            sys.stdout.write("\n")
            sys.stdout.flush()

        json.dump(
            run_pipeline(request, emit if request.get("streamResults") else None),
            sys.stdout,
            ensure_ascii=False,
        )
        sys.stdout.write("\n")
        return 0
    except Exception as exc:
        json.dump(
            {"status": "FAILED", "errorType": type(exc).__name__, "error": str(exc)},
            sys.stdout,
            ensure_ascii=False,
        )
        sys.stdout.write("\n")
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
