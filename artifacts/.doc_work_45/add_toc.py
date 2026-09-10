from __future__ import annotations

import argparse
import copy
import json
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Twips


SOURCE = Path(
    "/Users/sunnytearlie/Downloads/datagraph-bank/artifacts/"
    "《面向商业银行非法行为识别的知识模型构建技术方案》（第四、五章补充稿）.docx"
)
OUTPUT = Path(
    "/Users/sunnytearlie/Downloads/datagraph-bank/artifacts/"
    "《面向商业银行非法行为识别的知识模型构建技术方案》（第四、五章补充稿，含目录）.docx"
)


def set_font_properties(element, family: str, size_pt: float, bold: bool = False) -> None:
    rpr = element.get_or_add_rPr() if hasattr(element, "get_or_add_rPr") else element
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("ascii", "hAnsi", "eastAsia"):
        rfonts.set(qn(f"w:{attr}"), family)
    for tag in ("w:sz", "w:szCs"):
        node = rpr.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            rpr.append(node)
        node.set(qn("w:val"), str(round(size_pt * 2)))
    for tag in ("w:b", "w:bCs"):
        node = rpr.find(qn(tag))
        if bold:
            if node is None:
                node = OxmlElement(tag)
                rpr.append(node)
            node.set(qn("w:val"), "1")
        elif node is not None:
            rpr.remove(node)


def style_font(style, family: str, size_pt: float, bold: bool = False) -> None:
    style.font.name = family
    style.font.size = Pt(size_pt)
    style.font.bold = bold
    rpr = style.element.get_or_add_rPr()
    set_font_properties(rpr, family, size_pt, bold)


def ensure_toc_styles(doc: Document) -> None:
    specs = {
        "TOC": ("Heiti SC", 18, True, 0, 18),
        "TOC 1": ("Heiti SC", 14, True, 0, 7),
        "TOC 2": ("Songti SC", 14, False, 28, 6),
        "TOC 3": ("Songti SC", 12, False, 56, 5),
    }
    for name, (font, size, bold, left, after) in specs.items():
        try:
            style = doc.styles[name]
        except KeyError:
            style = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style_font(style, font, size, bold)
        fmt = style.paragraph_format
        fmt.left_indent = Pt(left)
        fmt.right_indent = Pt(0)
        fmt.first_line_indent = Pt(0)
        fmt.space_before = Pt(0)
        fmt.space_after = Pt(after)
        fmt.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        fmt.line_spacing = 1.15
        fmt.keep_together = True
        fmt.widow_control = True
        if name != "TOC":
            fmt.tab_stops.add_tab_stop(
                Twips(8296), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS
            )


def add_run_properties(run, family: str, size_pt: float, bold: bool = False) -> None:
    rpr = run.find(qn("w:rPr"))
    if rpr is None:
        rpr = OxmlElement("w:rPr")
        run.insert(0, rpr)
    set_font_properties(rpr, family, size_pt, bold)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "000000")
    rpr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "none")
    rpr.append(underline)


def make_text_run(text: str, family: str, size_pt: float, bold: bool = False):
    run = OxmlElement("w:r")
    add_run_properties(run, family, size_pt, bold)
    t = OxmlElement("w:t")
    if text.startswith(" ") or text.endswith(" "):
        t.set(qn("xml:space"), "preserve")
    t.text = text
    run.append(t)
    return run


def make_tab_run(family: str, size_pt: float, bold: bool = False):
    run = OxmlElement("w:r")
    add_run_properties(run, family, size_pt, bold)
    run.append(OxmlElement("w:tab"))
    return run


def make_field_runs(instruction: str, result: str, family: str, size_pt: float, bold: bool = False):
    runs = []
    for kind in ("begin",):
        run = OxmlElement("w:r")
        add_run_properties(run, family, size_pt, bold)
        fld = OxmlElement("w:fldChar")
        fld.set(qn("w:fldCharType"), kind)
        run.append(fld)
        runs.append(run)
    run = OxmlElement("w:r")
    add_run_properties(run, family, size_pt, bold)
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    run.append(instr)
    runs.append(run)
    run = OxmlElement("w:r")
    add_run_properties(run, family, size_pt, bold)
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    run.append(sep)
    runs.append(run)
    runs.append(make_text_run(result, family, size_pt, bold))
    run = OxmlElement("w:r")
    add_run_properties(run, family, size_pt, bold)
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run.append(end)
    runs.append(run)
    return runs


def make_field_open_runs(instruction: str, family: str, size_pt: float, bold: bool = False):
    begin_run = OxmlElement("w:r")
    add_run_properties(begin_run, family, size_pt, bold)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin.set(qn("w:dirty"), "true")
    begin_run.append(begin)
    instr_run = OxmlElement("w:r")
    add_run_properties(instr_run, family, size_pt, bold)
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    instr_run.append(instr)
    separate_run = OxmlElement("w:r")
    add_run_properties(separate_run, family, size_pt, bold)
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    separate_run.append(separate)
    return [begin_run, instr_run, separate_run]


def make_field_end_run(family: str, size_pt: float, bold: bool = False):
    run = OxmlElement("w:r")
    add_run_properties(run, family, size_pt, bold)
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run.append(end)
    return run


def heading_level(paragraph) -> int | None:
    name = paragraph.style.name
    if name == "Heading 1":
        return 1
    if name == "Heading 2":
        return 2
    if name == "Heading 3":
        return 3
    return None


def add_bookmark(paragraph, bookmark_id: int, bookmark_name: str) -> None:
    p = paragraph._p
    for old in p.findall(qn("w:bookmarkStart")):
        if old.get(qn("w:name"), "").startswith("_DatagraphToc"):
            p.remove(old)
    for old in p.findall(qn("w:bookmarkEnd")):
        if old.get(qn("w:id")) == str(bookmark_id):
            p.remove(old)
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bookmark_id))
    start.set(qn("w:name"), bookmark_name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bookmark_id))
    insert_at = 1 if p.pPr is not None else 0
    p.insert(insert_at, start)
    p.append(end)


def create_toc_title(doc: Document):
    p = doc.add_paragraph(style="TOC")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.page_break_before = True
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(20)
    run = p.add_run("目  录")
    run.font.name = "Heiti SC"
    run.font.size = Pt(18)
    run.bold = True
    set_font_properties(run._element, "Heiti SC", 18, True)
    return p


def create_toc_entry(
    doc: Document,
    text: str,
    level: int,
    bookmark: str,
    page: str,
    *,
    open_toc_field: bool = False,
    close_toc_field: bool = False,
):
    style_name = f"TOC {level}"
    p = doc.add_paragraph(style=style_name)
    p.paragraph_format.keep_together = True
    p.paragraph_format.widow_control = True
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.tab_stops.add_tab_stop(
        Twips(8296), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS
    )
    if level == 1:
        family, size, bold = "Heiti SC", 14, True
    elif level == 2:
        family, size, bold = "Songti SC", 14, False
    else:
        family, size, bold = "Songti SC", 12, False
    if open_toc_field:
        for run in make_field_open_runs(
            ' TOC \\o "1-3" \\h \\z \\u ', family, size, bold
        ):
            p._p.append(run)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("w:anchor"), bookmark)
    hyperlink.set(qn("w:history"), "1")
    hyperlink.append(make_text_run(text, family, size, bold))
    hyperlink.append(make_tab_run(family, size, bold))
    for run in make_field_runs(
        f" PAGEREF {bookmark} \\h ", page, family, size, bold
    ):
        hyperlink.append(run)
    p._p.append(hyperlink)
    if close_toc_field:
        p._p.append(make_field_end_run(family, size, bold))
    return p


def clone_section_break(doc: Document):
    paragraph = doc.add_paragraph()
    ppr = paragraph._p.get_or_add_pPr()
    sectpr = copy.deepcopy(doc.sections[-1]._sectPr)
    for child in list(sectpr):
        if child.tag in {
            qn("w:headerReference"),
            qn("w:footerReference"),
            qn("w:pgNumType"),
            qn("w:type"),
        }:
            sectpr.remove(child)
    section_type = OxmlElement("w:type")
    section_type.set(qn("w:val"), "nextPage")
    sectpr.insert(0, section_type)
    ppr.append(sectpr)
    return paragraph


def move_before(elements, anchor) -> None:
    parent = anchor.getparent()
    index = parent.index(anchor)
    for element in elements:
        parent.insert(index, element)
        index += 1


def replace_page_number_type(section, fmt: str, start: int) -> None:
    sectpr = section._sectPr
    old = sectpr.find(qn("w:pgNumType"))
    if old is not None:
        sectpr.remove(old)
    node = OxmlElement("w:pgNumType")
    node.set(qn("w:fmt"), fmt)
    node.set(qn("w:start"), str(start))
    sectpr.append(node)


def set_footer_page_number(section, cached_result: str) -> None:
    section.footer.is_linked_to_previous = False
    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.clear()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = Pt(0)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    for run in make_field_runs(
        " PAGE \\* MERGEFORMAT ", cached_result, "Times New Roman", 10.5, False
    ):
        paragraph._p.append(run)


def enable_field_updates(doc: Document) -> None:
    settings = doc.settings._element
    node = settings.find(qn("w:updateFields"))
    if node is None:
        node = OxmlElement("w:updateFields")
        settings.append(node)
    node.set(qn("w:val"), "true")


def build(page_map: dict[str, int]) -> None:
    doc = Document(SOURCE)
    ensure_toc_styles(doc)
    headings = [p for p in doc.paragraphs if heading_level(p) is not None]
    if not headings:
        raise SystemExit("No Heading 1-3 paragraphs found.")

    title = create_toc_title(doc)
    toc_paragraphs = [title]
    max_existing_id = 0
    for node in doc._element.findall(".//" + qn("w:bookmarkStart")):
        try:
            max_existing_id = max(max_existing_id, int(node.get(qn("w:id"), "0")))
        except ValueError:
            pass
    for index, heading in enumerate(headings, 1):
        level = heading_level(heading)
        bookmark_id = max_existing_id + index
        bookmark_name = f"_DatagraphToc{index:03d}"
        add_bookmark(heading, bookmark_id, bookmark_name)
        page = str(page_map.get(heading.text, 1))
        toc_paragraphs.append(
            create_toc_entry(
                doc,
                heading.text,
                level,
                bookmark_name,
                page,
                open_toc_field=index == 1,
                close_toc_field=index == len(headings),
            )
        )

    section_break = clone_section_break(doc)
    toc_paragraphs.append(section_break)
    first_heading = headings[0]
    first_heading.paragraph_format.page_break_before = False
    move_before([p._p for p in toc_paragraphs], first_heading._p)
    enable_field_updates(doc)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)

    doc = Document(OUTPUT)
    if len(doc.sections) != 2:
        raise SystemExit(f"Expected 2 sections after TOC insertion, got {len(doc.sections)}")
    replace_page_number_type(doc.sections[0], "upperRoman", 1)
    replace_page_number_type(doc.sections[1], "decimal", 1)
    set_footer_page_number(doc.sections[0], "I")
    set_footer_page_number(doc.sections[1], "1")
    enable_field_updates(doc)
    doc.save(OUTPUT)

    verify = Document(OUTPUT)
    verify_headings = [p.text for p in verify.paragraphs if heading_level(p) is not None]
    if verify_headings != [p.text for p in headings]:
        raise SystemExit("Heading sequence changed while inserting TOC.")
    if len(verify.sections) != 2:
        raise SystemExit("Section structure changed unexpectedly.")
    print(OUTPUT)
    print(f"toc_entries={len(headings)} sections={len(verify.sections)}")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--pages-json", type=Path)
    args = parser.parse_args()
    page_map: dict[str, int] = {}
    if args.pages_json:
        page_map = json.loads(args.pages_json.read_text(encoding="utf-8"))
    build(page_map)


if __name__ == "__main__":
    main()
