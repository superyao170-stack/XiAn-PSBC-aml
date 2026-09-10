from __future__ import annotations

import hashlib
import shutil
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Twips


REFERENCE = Path('/Users/sunnytearlie/Desktop/西安邮储/《面向商业银行非法行为识别的知识模型构建技术方案》.docx')
OUTPUT = Path('/Users/sunnytearlie/Downloads/datagraph-bank/artifacts/《面向商业银行非法行为识别的知识模型构建技术方案》（第四、五章补充稿）.docx')
EXPECTED_REFERENCE_SHA256 = '8b0ad607d5c96a3690e3d740916b790308822760ce61b63ff302d6f0f898a1f8'
SKILL_SCRIPTS = Path('/Users/sunnytearlie/.codex/plugins/cache/openai-primary-runtime/documents/26.826.12353/skills/documents/scripts')
sys.path.insert(0, str(SKILL_SCRIPTS))
from table_geometry import apply_table_geometry, column_widths_from_weights, section_content_width_dxa


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def set_run_font(run, east_asia: str = 'Songti SC', ascii_font: str = 'Times New Roman', size: float | None = None, bold: bool | None = None):
    # LibreOffice on macOS may route CJK glyphs through the ascii/hAnsi font
    # even when w:eastAsia is present. Use the installed CJK family for every
    # range so both Word and the verification renderer display Chinese text.
    run.font.name = east_asia
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:ascii'), east_asia)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:hAnsi'), east_asia)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def configure_styles(doc: Document) -> None:
    style_specs = {
        'Normal': ('Songti SC', 'Times New Roman', 14, False),
        'Heading 1': ('Heiti SC', 'Times New Roman', 16, True),
        'Heading 2': ('Heiti SC', 'Times New Roman', 15, True),
        'Heading 3': ('Heiti SC', 'Times New Roman', 14, True),
    }
    for name, (east_asia, ascii_font, size, bold) in style_specs.items():
        style = doc.styles[name]
        style.font.name = east_asia
        style.font.size = Pt(size)
        style.font.bold = bold
        rpr = style.element.get_or_add_rPr()
        rfonts = rpr.get_or_add_rFonts()
        rfonts.set(qn('w:ascii'), east_asia)
        rfonts.set(qn('w:hAnsi'), east_asia)
        rfonts.set(qn('w:eastAsia'), east_asia)
    normal = doc.styles['Normal'].paragraph_format
    normal.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.line_spacing = 1.5
    normal.first_line_indent = Pt(28)
    normal.space_before = Pt(0)
    normal.space_after = Pt(0)
    for name in ('Heading 1', 'Heading 2', 'Heading 3'):
        fmt = doc.styles[name].paragraph_format
        fmt.first_line_indent = Pt(0)
        fmt.keep_with_next = True
        fmt.keep_together = True
    doc.styles['Heading 1'].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.styles['Heading 1'].paragraph_format.space_before = Pt(6)
    doc.styles['Heading 1'].paragraph_format.space_after = Pt(10)
    doc.styles['Heading 2'].paragraph_format.space_before = Pt(8)
    doc.styles['Heading 2'].paragraph_format.space_after = Pt(4)
    doc.styles['Heading 3'].paragraph_format.space_before = Pt(5)
    doc.styles['Heading 3'].paragraph_format.space_after = Pt(2)

    # Convert the reference's Windows-only Chinese font names to installed
    # macOS equivalents while keeping the same serif/sans roles.
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            east_asia = run._element.get_or_add_rPr().get_or_add_rFonts().get(qn('w:eastAsia'))
            if east_asia in {'黑体', 'SimHei'} or run.font.name in {'黑体', 'SimHei'}:
                set_run_font(run, 'Heiti SC', 'Times New Roman')
            elif east_asia in {'宋体', 'SimSun'} or run.font.name in {'宋体', 'SimSun'}:
                set_run_font(run, 'Songti SC', 'Times New Roman')


def add_heading(doc: Document, text: str, level: int, page_break_before: bool = False):
    if page_break_before:
        doc.add_page_break()
    p = doc.add_paragraph(style=f'Heading {level}')
    p.paragraph_format.first_line_indent = Pt(0)
    if level == 1:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    set_run_font(run, 'Heiti SC', 'Times New Roman', {1: 16, 2: 15, 3: 14}[level], True)
    return p


def add_body(doc: Document, text: str, *, bold_lead: str | None = None):
    p = doc.add_paragraph(style='Normal')
    p.paragraph_format.first_line_indent = Pt(28)
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if bold_lead and text.startswith(bold_lead):
        first = p.add_run(bold_lead)
        set_run_font(first, 'Heiti SC', 'Times New Roman', 14, True)
        rest = p.add_run(text[len(bold_lead):])
        set_run_font(rest, 'Songti SC', 'Times New Roman', 14, False)
    else:
        run = p.add_run(text)
        set_run_font(run, 'Songti SC', 'Times New Roman', 14, False)
    return p


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tc_pr.append(shd)
    shd.set(qn('w:fill'), fill)


def set_cell_no_wrap(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    if tc_pr.find(qn('w:noWrap')) is None:
        tc_pr.append(OxmlElement('w:noWrap'))


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement('w:tblHeader')
    header.set(qn('w:val'), 'true')
    tr_pr.append(header)


def set_row_cant_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn('w:cantSplit')) is None:
        tr_pr.append(OxmlElement('w:cantSplit'))


def add_table(doc: Document, caption: str, headers: list[str], rows: list[list[str]], weights: list[float], *, font_size: float = 10.5):
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.keep_with_next = True
    cap.paragraph_format.space_before = Pt(5)
    cap.paragraph_format.space_after = Pt(3)
    cap_run = cap.add_run(caption)
    set_run_font(cap_run, 'Heiti SC', 'Times New Roman', 11, True)

    table = doc.add_table(rows=1, cols=len(headers))
    table.style = 'Table Grid'
    header_row = table.rows[0]
    set_repeat_table_header(header_row)
    set_row_cant_split(header_row)
    for idx, text in enumerate(headers):
        cell = header_row.cells[idx]
        set_cell_shading(cell, 'E7E6E6')
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Pt(0)
        p.paragraph_format.line_spacing = 1.15
        r = p.add_run(text)
        set_run_font(r, 'Heiti SC', 'Times New Roman', font_size, True)
    for row_data in rows:
        row = table.add_row()
        set_row_cant_split(row)
        for idx, text in enumerate(row_data):
            cell = row.cells[idx]
            if idx == 0 and (headers[0] == '业务码' or (headers[0] == '方法' and len(str(text)) <= 6)):
                set_cell_no_wrap(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx == 0 and len(headers) > 2 else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(str(text))
            if idx == 0 and headers[0] == '方法':
                cell_font_size = 8
            elif idx == 0 and headers[0] == '业务码':
                cell_font_size = 9
            else:
                cell_font_size = font_size
            set_run_font(r, 'Songti SC', 'Times New Roman', cell_font_size, False)
    total_width = section_content_width_dxa(doc.sections[-1]) - 120
    widths = column_widths_from_weights(weights, total_width)
    apply_table_geometry(table, widths, table_width_dxa=total_width, indent_dxa=120,
                         cell_margins_dxa={'top': 90, 'bottom': 90, 'start': 120, 'end': 120})
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(0)
    after.paragraph_format.first_line_indent = Pt(0)
    return table


def add_code(doc: Document, text: str):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Pt(18)
    p.paragraph_format.right_indent = Pt(12)
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(3)
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), 'F2F2F2')
    p_pr.append(shd)
    for idx, line in enumerate(text.splitlines()):
        if idx:
            p.add_run().add_break()
        r = p.add_run(line)
        set_run_font(r, 'Songti SC', 'Menlo', 9.5, False)
    return p


def chapter_four(doc: Document) -> None:
    add_heading(doc, '四、多源数据流处理引擎与案例智能处理功能说明', 1, page_break_before=True)
    add_body(doc, '本章面向系统使用和业务验收，说明案例数据从接入、校验、入库到智能处理、人工复核和图谱展示的完整流程。系统当前支持反洗钱与反欺诈两类监管场景，并针对新增案例和历史案例设置不同的处理路线。在新增案例路线中，可疑报告生成、五层框架抽取和相似案例检索构成三个连续的核心功能；在历史案例路线中，系统复用已有分析文本，优先完成框架抽取和图谱入库，以提高批量案例库建设效率。')
    add_body(doc, '功能实现以第一章定义的案例基本信息层、实体层、事件层、关系层和证据层为统一知识结构，以第二章本体和关系约束为语义边界，以第三章分类分级指标为风险解释基础。每一处理阶段都保存输入、输出、状态和操作记录，使自动分析结果能够被查询、复核和重放。')

    add_heading(doc, '4.1 引擎总体架构与处理流程', 2)
    add_heading(doc, '4.1.1 引擎定位与功能目标', 3)
    add_body(doc, '多源数据流处理引擎是案例识别系统的业务入口和流程编排中心。其主要目标不是对上传文件进行简单存储，而是把不同场景、不同来源和不同结构的案例材料转换为可以持续计算、比较和图谱化表达的标准案例知识。引擎统一组织文件校验、任务创建、分析Worker调用、阶段状态迁移、结果入库和异常反馈，避免各算法模块直接面对不一致的原始数据。')
    add_body(doc, '系统同时强调处理可控和结果可追溯。上传数据保留来源和内容指纹，案例处理池记录每个阶段的产物，分析任务记录运行进度和错误信息，人工修改和审批形成审计轨迹。由此，分析人员既能够按流程完成案例处理，也能够从最终结论反向查看所依据的数据、报告、知识对象和相似案例。')

    add_heading(doc, '4.1.2 系统主要功能组成', 3)
    add_body(doc, '系统功能由案例上传、数据预处理、案例智能处理、复核审批、案例管理和图谱分析六部分组成。案例上传负责接收单案例JSON文件或批量CSV/XLSX文件；数据预处理负责检查结构、字段、案例编号和场景一致性；案例智能处理依次完成可疑报告生成、五层框架抽取和相似案例检索；复核审批负责风险定级和结果确认；案例管理提供查询、详情和处理记录；图谱分析负责全景展示、关联线索分析和隐蔽风险挖掘。')
    add_body(doc, '上述功能通过统一案例编号、银行编码、场景编码和任务编号衔接。PostgreSQL保存业务状态、结构化知识成果和审计记录，TuGraph保存可遍历的案例图结构，Python Worker承载报告生成、框架抽取和相似度计算，前端工作台负责文件上传、任务控制、结果复核和图谱交互。')

    add_heading(doc, '4.1.3 案例数据端到端处理流程', 3)
    add_body(doc, '新增案例的主流程为：案例上传 -> 数据校验 -> 案例与任务入库 -> 可疑报告生成 -> 五层框架抽取 -> 相似案例检索 -> 风险等级推荐 -> 人工复核审批 -> 案例库激活 -> 关系图谱展示。系统只允许案例在满足当前阶段前置条件时进入下一阶段，例如未生成可疑报告的新增案例不能直接执行框架抽取，未完成框架抽取的案例不能执行相似匹配。')
    add_body(doc, '处理状态同时写入案例表和案例处理池。页面按“待生成报告、待框架抽取、待相似匹配、待复核审批”展示不同队列，操作完成后由服务端原子更新阶段状态和阶段产物。失败案例保持在原业务队列并记录错误原因，可在修正输入或运行条件后重新执行。')
    add_table(doc, '表4-1 新增案例处理阶段及主要产物',
              ['处理阶段', '系统状态', '主要输入', '主要输出'], [
                  ['数据接入', '上传任务', '案例JSON或CSV/XLSX', '上传凭据、任务记录、原始载荷'],
                  ['可疑报告', 'PENDING_REPORT', '基础信息、客户及场景数据', '风险事件链、分析文本、可疑报告'],
                  ['框架抽取', 'PENDING_EXTRACTION', '原始载荷与可疑报告', '五层知识结果、图谱快照'],
                  ['相似匹配', 'PENDING_SIMILARITY', '五层知识结果与已审核案例库', 'Top-5相似案例、得分及匹配子图'],
                  ['复核审批', 'PENDING_APPROVAL', '报告、框架、相似结果和推荐等级', '最终风险等级、人工排序、审批记录'],
                  ['成果使用', 'APPROVED', '已审核案例及图谱投影', '案例库条目、全景图谱与分析入口'],
              ], [1.2, 1.4, 3.0, 3.2])

    add_heading(doc, '4.1.4 新增案例与历史案例分流机制', 3)
    add_body(doc, '新增案例用于处理尚未形成完整分析结论的材料。上传阶段只保存基础事实，案例进入待生成报告队列，由系统生成或补充分析文本，再依次进入框架抽取、相似匹配和人工复核。该路线强调机器分析与人工确认的逐级收敛，任何报告修改都会使依赖它的框架、图谱快照和相似结果失效，防止旧结果继续参与审批。')
    add_body(doc, '历史案例用于建设可复用的已知案例库。历史反洗钱案例必须提供analysis_texts，历史反欺诈案例必须提供text_analysis；系统不重复生成报告，而是直接执行五层框架抽取。若基本信息中的风险等级可以映射为现行低、中、高三级，系统自动完成审核并进入图谱；风险等级缺失或无法识别时，案例转入人工复核队列。历史案例在本次接入中不执行相似度匹配，相似检索由后续新增案例显式触发。')

    add_heading(doc, '4.2 多源案例数据接入与预处理', 2)
    add_heading(doc, '4.2.1 统一数据接入方式', 3)
    add_body(doc, '系统通过统一案例上传工作台接收反洗钱和反欺诈案例。用户首先选择处理场景、案例来源和处理方式，再上传对应文件。上传接口只负责接收和校验材料并生成上传凭据，分析任务通过独立接口创建和启动，使大文件上传与长时间分析相互解耦。任务列表展示任务编号、任务名称、银行、场景、状态、进度和创建时间，并提供查看结果、停止、重试和删除等操作。')
    add_body(doc, '单案例模式适用于调试、补录和重点案例处理，以分角色JSON文件表达案例基本信息、客户、账户、设备或其他实体。批处理模式适用于历史库建设和批量新增案例识别，使用CSV或XLSX文件，一行对应一个案例，复杂字段以JSON字符串写入单元格。批处理文件最多包含500个案例，系统拒绝重复case_id、空文件、公式单元格和不符合场景的表头。')

    add_heading(doc, '4.2.2 单案例与批量案例上传', 3)
    add_body(doc, '单案例上传采用明确的角色文件边界。basic_info.json顶层必须为对象并包含非空case_id；customers.json顶层必须为非空数组，每个客户必须包含entity_id。反洗钱场景可选accounts.json和other_entities.json；反欺诈场景要求accounts.json和devices.json，并通过设备号或device_id识别设备。单案例文件及全部文件合计不得超过100MB。')
    add_body(doc, '批处理上传只接受一个CSV或XLSX文件。反洗钱新增案例表头为basic_info、customers、transaction_features，历史案例表头至少为basic_info、customers、analysis_texts；反欺诈新增案例增加accounts、devices和event_chain，历史案例以text_analysis替代event_chain。系统按行解析、验证和拆分案例，并为每个案例建立独立的结果和状态记录。')
    add_table(doc, '表4-2 主要案例输入模式',
              ['场景与来源', '单案例主要文件', '批处理主要字段', '后续路线'], [
                  ['反洗钱新增案例', 'basic_info、customers；accounts和other_entities可选', 'basic_info、customers、transaction_features', '生成报告后进行框架抽取和相似匹配'],
                  ['反洗钱历史案例', 'basic_info、customers、analysis_texts；其他实体可选', 'basic_info、customers、analysis_texts', '复用报告并直接进行框架抽取'],
                  ['反欺诈新增案例', 'basic_info、customers、accounts、devices、event_chain', 'basic_info、customers、accounts、devices、event_chain', '规则生成风险事件链后生成报告'],
                  ['反欺诈历史案例', 'basic_info、customers、accounts、devices、text_analysis', 'basic_info、customers、accounts、devices、text_analysis', '复用报告并直接进行框架抽取'],
              ], [1.4, 3.0, 3.0, 2.4])

    add_heading(doc, '4.2.3 反洗钱与反欺诈数据要求', 3)
    add_body(doc, '反洗钱与反欺诈使用相同的案例编号、银行范围和五层输出契约，但输入事实和规则库相互隔离。反洗钱重点使用客户、账户、交易特征和已有分析文本；反欺诈增加设备和渠道信息，并使用独立的分渠道风险规则库。系统在上传阶段识别明显的跨场景文件，例如选择反洗钱场景却上传反欺诈批处理表头时直接阻断，避免错误规则和提示词作用于不匹配的数据。')
    add_body(doc, '反欺诈的客户、账户、设备和案例基本信息采用确定性直映射，模型主要负责风险事件和关系的抽取与组织。规则命中表示待核验风险线索，不直接构成诈骗或违法认定；报告和框架结果仍需结合原始材料、事件证据和人工复核形成最终判断。')

    add_heading(doc, '4.2.4 数据校验、清洗与标准化', 3)
    add_body(doc, '预处理首先检查文件扩展名、文件大小、JSON顶层类型、必填文件和必填字段，再检查案例编号、客户编号和场景模式的一致性。CSV/XLSX解析采用固定表头，禁止额外字段和公式，避免在服务端执行不受控表达式。对文本采用UTF-8兼容读取，对空白值进行裁剪，对风险等级、模式代码和场景代码进行标准化。')
    add_body(doc, '案例编号是贯穿业务库、处理池、图谱和案例库的主标识。系统拒绝跨银行覆盖同一案例，批处理内部也不允许重复case_id。相同银行内的重复识别被视为对同一案例的权威重算，新的框架、事件、图谱快照和处理状态替换旧结果，同时保留任务关系和审计记录。')
    add_body(doc, '系统还生成源数据内容指纹和图谱快照哈希，用于判断输入或结果是否发生变化。字段清洗只改变可安全确定的格式，不凭空补充业务事实；无法判断的风险等级、时间、金额或身份信息保留为空或待核验状态，并在后续复核中提示。')

    add_heading(doc, '4.2.5 案例入库与处理任务创建', 3)
    add_body(doc, '上传校验通过后，服务端在受控上传目录保存文件并返回uploadToken。用户填写任务名称、银行和场景后创建分析任务，任务输入参数记录workflow、processingMode、recognitionMode、uploadToken和caseCount。任务由后端线程池异步执行，默认并发度为5，队列容量为1000；页面通过任务详情接口轮询进度，不需要保持长连接等待。')
    add_body(doc, '新增案例在初始入库时写入案例表和案例处理池，状态为PENDING_REPORT；历史案例由Worker直接产出框架结果，根据基础信息风险等级进入APPROVED或PENDING_APPROVAL。任务停止、失败和重试不会绕过阶段校验，删除任务时系统按任务与案例的归属关系清理独占成果，同时保留输入批次和原始文件。')

    add_heading(doc, '4.3 模块一：可疑报告智能生成', 2)
    add_heading(doc, '4.3.1 功能定位与输入数据', 3)
    add_body(doc, '可疑报告生成是新增案例进入知识抽取前的第一项核心功能，用于把分散的基础信息、客户信息、账户信息、交易特征、设备信息和事件链整理为可阅读、可复核的案例分析文本。该功能只处理PENDING_REPORT状态的新增案例；历史案例已有分析文本，不重复调用报告生成流程。')
    add_body(doc, '服务端从案例处理池读取source_payload，在受控临时目录恢复各角色JSON文件，并将场景代码、识别模式和目标阶段传给对应Worker。反洗钱与反欺诈使用平行Worker目录和独立提示词，避免业务规则交叉；共同的执行时限由后端统一控制，超时或Worker返回失败时保留原状态并记录last_error。')

    add_heading(doc, '4.3.2 风险事件链生成', 3)
    add_body(doc, '反欺诈新增案例首先依据basic_info中的渠道选择分渠道规则，对event_chain中的原始事件进行匹配，形成可审计的风险事件链。规则库只负责识别事件是否满足已登记的渠道特征，并记录规则编号、事件引用和命中依据；语义风险事件知识库则负责为事件提供统一名称、类别和说明，两类知识分开治理。')
    add_body(doc, '风险事件链按资金行为、交易对手、账户操作、设备与渠道、外部权威线索和时间顺序组织。每个进入报告的风险描述必须引用原始事件或确定性派生事实，无法获得证据的内容标记为不可评估。反洗钱场景不套用反欺诈渠道规则，而是直接依据客户、账户和交易特征构建报告分析上下文。')

    add_heading(doc, '4.3.3 案例分析与可疑报告生成', 3)
    add_body(doc, '报告生成采用分阶段、分角色的并行编排。反洗钱流程分别分析主体情况、客户画像与交易特征、发现过程、资金链路、案例主体、关联账户和交易对手等内容；反欺诈流程并行分析客户/账户/设备概览、资金链、案例主体、资金与操作、交易对手与外部线索，再生成客户主体综合分析和案例综合分析。独立节点只接收与本段相关的事实，降低长上下文中事实混用的风险。')
    add_body(doc, '各节点返回结构化段落及其事实引用，工作流检查段落编号、内容完整性、证据引用和长度约束。初稿完成后，统一复核节点检查事实矛盾、无依据判断、遗漏的重要风险和段落重复；仅对存在错误的目标段落执行定向重写，并设置最大重写轮次。若复核后仍存在关键错误，流程返回失败而不发布报告。')
    add_body(doc, '最终产物包含用于页面展示的analysisText、保留分段结构的analysisTexts、风险事件链、复核结果和必要的执行元数据。报告文字强调“异常事实、风险解释、待核验事项”的区分，不以模型生成内容替代法律性质认定或人工审核。')

    add_heading(doc, '4.3.4 报告查看、编辑与保存', 3)
    add_body(doc, '处理完成后，用户可在案例处理详情中同时查看原始材料、风险事件链和可疑报告。系统允许对analysisText进行人工修订，保存时将来源标记为USER_EDITED_ANALYSIS_TEXT，并保留既有风险事件链。修改后的报告重新进入PENDING_EXTRACTION，原框架结果、图谱快照、相似结果、推荐得分和旧相似排序被清空。')
    add_body(doc, '该设计保证下游结果始终基于当前有效报告。人工编辑不直接修改原始业务数据，也不改变风险事件链引用；若事实本身有误，应回到数据接入环节重新上传或重跑案例，而不是只修改结论性文字。')

    add_heading(doc, '4.3.5 处理结果与任务状态流转', 3)
    add_body(doc, '可疑报告生成成功后，案例由PENDING_REPORT迁移为PENDING_EXTRACTION，并在case_processing_audit中记录原状态、目标状态、动作、报告快照、操作人和时间。批量处理逐案执行并汇总成功、失败数量，单个案例失败不阻断其他案例。')
    add_body(doc, '失败通常来自输入缺失、模型服务不可用、响应格式错误、事实引用校验不通过或执行超时。页面展示错误信息并允许重试；在重试成功前，案例仍停留在待生成报告队列，不会进入框架抽取。')

    add_heading(doc, '4.4 模块二：知识驱动的五层框架抽取', 2)
    add_heading(doc, '4.4.1 功能定位与抽取流程', 3)
    add_body(doc, '五层框架抽取是将可疑报告和原始案例材料转换为标准知识对象的核心环节。新增案例以已确认的可疑报告作为语义输入，历史案例直接复用上传的分析文本。Worker在统一工作流中完成数据读取、实体和事件抽取、分层关系生成、证据绑定、语义校验和图谱快照组装，输出结构固定的final_case数据。')
    add_body(doc, '抽取遵循“确定性数据优先、模型抽取受约束、结果必须可验证”的原则。基础信息、客户、账户和设备等已有结构字段尽量直接映射；模型重点处理报告中的事件、关系和证据指向；所有关系必须满足允许的源类型、目标类型和方向语义，不符合契约的候选结果被过滤或进入重试。')

    add_heading(doc, '4.4.2 案例基本信息与实体抽取', 3)
    add_body(doc, '案例基本信息层从basic_info中读取案例编号、名称、来源编号、场景、业务领域、案例类型、触发点、报告时间、业务状态、风险等级、涉嫌类型和处置措施等信息。案例编号作为稳定主键，案例名称用于业务展示，两者分开管理。缺失字段不由模型随意补写，必要时保留为空并在案例详情中由人工维护。')
    add_body(doc, '实体层包括客户、账户、设备和其他实体。系统根据entity_id、账户标识、设备标识和案例范围生成稳定知识标识，对同名异体、异名同体和跨文件重复对象执行归一化。反欺诈场景的客户、账户和设备采用直映射，反洗钱场景可在原有客户基础上补充账户和其他实体。实体输出保留类型、名称、关键属性、来源和置信度。')

    add_heading(doc, '4.4.3 事件、关系与证据抽取', 3)
    add_body(doc, '事件层从报告和风险事件链中识别交易、账户操作、身份变化、外部处置及其他风险行为，记录事件编号、名称、类型、时间、参与主体、动作对象、金额、渠道和证据引用。事件名称与系统事件元数据进行匹配，优先使用已有标准事件；未完全匹配的结果保留原始表述和候选标准编码，供治理人员复核。')
    add_body(doc, '关系层按照确定性结构关系、实体与实体关系、实体与事件关系、事件与事件关系逐层生成。典型关系包括包含、来源、持有、参与、涉及、顺承、上下位和应对关系。每层只接受契约允许的节点组合，最终按“关系类型、源节点、目标节点”去重，并记录关系来源和确定性。')
    add_body(doc, '证据层把原始文件、事实片段、事件引用和模型结论连接起来。证据对象保存来源类型、来源标识、摘要、校验信息和证明对象；报告中的每个非不可评估判断应能回溯到一个或多个证据引用。对于仅由规则或模型推断的关系，系统标记为派生关系并保留生成依据，不与原始事实关系混淆。')

    add_heading(doc, '4.4.4 实体归一与事件语义标准化', 3)
    add_body(doc, '实体归一首先使用业务标识进行精确匹配，在标识不足时结合名称、证件、联系方式、地址、账户归属和设备特征形成候选匹配。系统不在证据不足时强制合并实体，而是保留候选关系和置信度，避免错误合并扩散到事件链和跨案例分析。')
    add_body(doc, '事件语义标准化由事件字典解析器和事件语义增强服务共同完成。系统按照AML或ANTI_FRAUD场景加载对应事件知识库，将事件候选与统一的id、name、category和rule定义对齐，并保存知识库版本。CoreNLP可用于中文指代消解，LLM用于受约束的语义提取；服务不可用时按配置进行重试或采用已有结构信息继续处理。')

    add_heading(doc, '4.4.5 五层知识结果与图谱快照生成', 3)
    add_body(doc, '五层结果组装完成后，系统执行结构完整性、节点引用、关系方向、重复对象、证据引用和案例范围校验。通过校验的结果写入案例库文档，并生成包含节点数、关系数和完整图数据的graphSnapshot。新增案例的案例库记录在复核通过前保持非活动状态，避免未经审核的内容参与相似检索；历史案例按照识别结果进入自动审核或人工复核。')
    add_body(doc, '框架抽取成功后，案例进入PENDING_SIMILARITY。与此同时，事件等关系型对象写入PostgreSQL，Java侧图谱写入服务在合同投影校验通过后将案例、客户、账户、交易、事件和证据投影到TuGraph。图谱快照用于结果展示和故障回退，TuGraph用于全景查询和关系遍历。')

    add_heading(doc, '4.5 模块三：语义增强的相似案例检索', 2)
    add_heading(doc, '4.5.1 功能定位与历史案例库', 3)
    add_body(doc, '相似案例检索用于回答“当前新增案例与哪些已审核案例在事件语义和关系结构上相近”。检索输入是当前案例的五层框架结果，候选集来自同一银行、同一监管场景、状态为APPROVED且案例库记录为ACTIVE的历史案例。系统根据案例库数量和更新时间生成缓存键，只在案例库发生变化时重建历史快照，降低重复读取成本。')
    add_body(doc, '待查询案例不会与自身比较，未审核、已删除或跨银行案例不会进入候选集。历史案例库为空时系统返回空结果并说明未执行配对，不制造虚假相似案例；该案例仍可进入复核，由人工依据报告和框架结果定级。')

    add_heading(doc, '4.5.2 候选案例召回与语义匹配', 3)
    add_body(doc, '算法首先根据事件类型向量进行粗召回，筛除在行为构成上明显无关的案例。对召回候选中的事件描述，使用bge-small-zh-v1.5生成中文句向量并计算余弦相似度；达到召回阈值的事件对再由bge-reranker-v2-m3进行交叉编码重排，为后续节点对齐提供更可靠的语义代价。')
    add_body(doc, '语义模型只用于候选召回和事件描述对齐，不直接决定最终案例排序。系统记录向量相似度、重排得分、模型来源和召回阈值，便于解释某一事件对为什么被视为候选匹配。')

    add_heading(doc, '4.5.3 案例图结构相似度计算', 3)
    add_body(doc, '候选召回后，系统把查询案例和历史案例表示为异构案例图，节点包含案例、实体和事件，边表示参与、持有、包含、时序及其他业务关系。近似图编辑距离算法计算从一张案例图转换为另一张案例图所需的节点插入、删除、替换以及边变化成本，并结合事件语义相似度优化节点对齐。')
    add_body(doc, '原始图编辑距离按查询图完整删除成本与候选图完整插入成本进行归一化，最终相似度为1减去归一化图编辑距离。该指标同时反映事件内容、参与主体和关系结构差异，避免仅凭关键词重合判断案例相似。系统保留匹配节点、匹配关系、事件相似点和差异说明，用于页面对比。')

    add_heading(doc, '4.5.4 相似案例综合排序', 3)
    add_body(doc, '系统按归一化图编辑距离从小到大排序，并将其转换为0至1范围的相似度，默认返回前5个历史案例。结果记录算法排序algorithmRank、最终排序finalRank、相似度similarity、归一化权重normalizedWeight和是否人工调整manuallyReordered。当前版本以归一化近似图编辑距离作为最终排序指标，语义模型用于提高召回和对齐质量。')
    add_table(doc, '表4-3 相似案例检索的主要处理环节',
              ['处理环节', '主要方法', '作用', '输出'], [
                  ['候选召回', '事件类型向量', '缩小历史案例搜索范围', 'GED候选集合'],
                  ['语义对齐', 'BGE向量与交叉编码器', '识别描述不同但语义相近的事件', '事件对齐分数'],
                  ['结构匹配', '异构图近似GED', '比较节点和关系结构差异', '归一化GED及匹配子图'],
                  ['结果排序', '1 - normalized GED', '形成稳定、可解释的Top-5列表', '排名、相似度和权重'],
              ], [1.3, 2.2, 3.0, 2.5])

    add_heading(doc, '4.5.5 匹配结果对比与人工调整', 3)
    add_body(doc, '页面以排序列表和匹配图对比展示相似案例。复核人员可以查看双方基本信息、事件之间的相似点、匹配子图和主要差异，并在审批前调整相似案例顺序。人工排序必须包含算法返回的全部案例且不能重复，系统按调序后的名次重新计算调和归一化权重，同时保留算法原始名次。')
    add_body(doc, '当语义模型加载失败或超过独立时限时，系统自动降级为事件类型向量召回加结构化GED匹配，并在结果中记录degraded和degradationReason。若降级流程仍超时，则案例保留在PENDING_SIMILARITY并返回错误，避免以不完整排序进入审批。')

    add_heading(doc, '4.6 案例复核与风险定级', 2)
    add_heading(doc, '4.6.1 复核材料汇总展示', 3)
    add_body(doc, '复核页面将原始材料、风险事件链、可疑报告、五层框架结果、图谱快照、相似案例及推荐等级集中展示。复核人员不需要在多个系统间拼接证据，可以从风险等级反向查看事件、关系和相似案例贡献，并识别模型结论与原始事实是否一致。')

    add_heading(doc, '4.6.2 系统推荐风险等级', 3)
    add_body(doc, '当前系统使用可解释的确定性评分生成默认建议，评分由事件严重程度、风险类型关键词、关系网络复杂度、主体暴露程度、最高相似案例风险和来源风险提示组成，总分限制在0至100。0至39.99建议为低风险，40至69.99建议为中风险，70至100建议为高风险，并保存各分项得分、事实数量、阈值和模型版本。推荐结果只用于辅助复核，最终等级由审批人员确认。')
    add_body(doc, '第三章提出的分类分级体系具有扩展性，而当前软件复核接口落地为低、中、高三级。部署机构如需保留“较高风险”等更细等级，应通过统一的指标版本和接口契约扩展，不应在页面、数据库和算法中分别定义不同口径。')

    add_heading(doc, '4.6.3 人工复核与等级核定', 3)
    add_body(doc, '审批人员可以结合可疑报告、框架事实、证据充分度和相似案例重新核定风险等级，并调整相似案例最终顺序。审批通过后，案例状态更新为APPROVED，审批人和时间写入案例与处理池，结构化案例库记录激活，后续新增案例即可把它作为相似检索候选。')

    add_heading(doc, '4.6.4 历史案例自动审核机制', 3)
    add_body(doc, '历史案例的风险等级以final_case.basic_info.risk_level为唯一自动审核依据。系统支持01/低/低风险/LOW、02/中/中风险/中等风险/重点可疑/MEDIUM、03/高/高风险/HIGH等常见表达。成功映射时由analysis-worker记录自动审批；无法映射时转入PENDING_APPROVAL，禁止根据文本关键词或相似案例替代基础信息完成自动定级。')

    add_heading(doc, '4.6.5 审核结果与案例库更新', 3)
    add_body(doc, '审核结果形成案例状态、风险等级、审批人、相似案例最终排序和处理审计。报告、框架或基本事实后续发生修改时，应重新执行受影响的阶段并生成新的结果快照；案例库使用内容SHA-256判断版本变化。通过“机器推荐、人工确认、结果激活、后续复用”的闭环，系统逐步积累经过审核的高质量案例知识。')

    add_heading(doc, '4.7 非法行为关系图谱构建与展示', 2)
    add_heading(doc, '4.7.1 五层知识对象的图谱映射', 3)
    add_body(doc, '图谱投影以案例为组织节点，将客户、账户、交易、事件和证据映射为不同类型节点，并使用包含事件、包含证据、参与、持有账户、交易和证据支持等关系连接。知识解释、侧面事实、绑定和叙事对象保留在PostgreSQL，不全部投影到TuGraph，避免图数据库承担不适合遍历的辅助对象。')

    add_heading(doc, '4.7.2 图谱数据写入与更新', 3)
    add_body(doc, 'Java侧TuGraphStructuredWriter是图谱写入的唯一权威边界。写入前先捕获并校验关系型投影合同，确认节点标识、关系类型和案例范围合法；随后在BankGraph图中创建或复用Schema，按案例清理旧投影并写入新节点和关系。重复识别同一案例时采用重建方式，防止旧事件或旧关系残留。')
    add_body(doc, '系统同时保存WORKER_JSON类型的图谱快照及其哈希、节点数、关系数和完成状态。TuGraph不可用或案例图为空时，可使用快照或后端回退接口生成最小案例子图，以保证结果仍可查看；回退数据不得覆盖权威图谱。')

    add_heading(doc, '4.7.3 全景图谱展示', 3)
    add_body(doc, '全景图谱以已纳入业务范围的案例为入口，支持按银行、监管场景、案例和风险等级筛选。页面合并案例子图以及已审核案例之间的相似关系，显示当前范围内节点数和关系数，并提供关系探索布局。用户可以缩放、拖动画布、拖拽节点、适应画布和重新布局。')

    add_heading(doc, '4.7.4 节点、关系与证据详情查看', 3)
    add_body(doc, '点击节点后，系统展示节点类型、业务名称、风险等级、时间、金额、来源和案例归属等属性；点击关系可查看关系类型、方向语义、来源和确定性。对事件或知识解释对象，页面继续提供核心事件链、事项说明、声明和证据引用，使图形关系与可核验事实保持一致。')

    add_heading(doc, '4.7.5 相似案例与跨案例关联展示', 3)
    add_body(doc, '相似案例排序在图谱中表现为SIMILAR_CASE关系，关系属性包含最终名次、相似度、归一化权重和人工调整标记。跨案例视图帮助分析人员识别重复主体、共同账户、相似事件链和复现行为模式，但跨案例关联只提供风险线索，不自动改变单个案例的审批结论。')

    add_heading(doc, '4.7.6 关联线索分析与隐蔽风险挖掘', 3)
    add_body(doc, '全景图谱可继续进入关联线索分析和隐蔽风险挖掘工作台。关联线索分析包括事件链簇、行为伴生矩阵、风险传导与干预、事理模式增量学习；隐蔽风险挖掘包括元路径偏离、深层时序异常、局部超图分析和级联推理。各分析任务保存输入范围、参数、运行结果和分析资产，并对需要进入知识库的候选结果执行人工审核。')

    add_heading(doc, '4.8 本章小结', 2)
    add_body(doc, '本章说明了系统从多源案例数据接入到关系图谱展示的完整功能链。统一接入和预处理保证数据结构可靠；可疑报告生成将原始事实组织为可审阅文本；五层框架抽取把文本和结构数据转换为标准知识对象；相似案例检索通过语义召回与图结构匹配提供历史参照；复核审批确保最终风险等级和案例库内容由人工确认；图谱展示则把案例内部和跨案例关系转化为可探索、可追溯的知识网络。')


def chapter_five(doc: Document) -> None:
    add_heading(doc, '五、案例智能处理系统技术框架与API手册', 1, page_break_before=True)
    add_body(doc, '本章介绍当前项目的技术架构、组件接口和REST API。接口说明以仓库中的Spring Boot控制器、前端API适配、Worker调用协议和部署配置为依据，基础路径统一为/api/v1。系统版本为1.5.0，后端默认监听3030端口，前端通过同源反向代理访问业务接口。')

    add_heading(doc, '5.1 系统架构与技术框架', 2)
    add_heading(doc, '5.1.1 总体技术架构', 3)
    add_body(doc, '系统采用前后端分离、业务服务与算法Worker解耦、关系数据库与图数据库协同的架构。浏览器端负责交互和可视化；Spring Boot后端负责身份认证、银行数据范围、业务流程、状态机、审计和持久化；Python Worker负责报告生成、五层抽取和相似度计算；PostgreSQL保存业务事实、任务、案例库和分析资产；TuGraph保存面向遍历的案例图；CoreNLP和BGE模型服务提供中文语义处理能力。')
    add_body(doc, '组件之间使用明确接口连接。前端与后端采用JSON或multipart/form-data REST接口；后端与案例Worker采用JSON-over-stdio子进程协议；后端与风险分析服务采用HTTP JSON；后端与TuGraph通过Bolt/Cypher连接；Worker与CoreNLP及模型服务通过HTTP调用。各组件能够独立部署和扩缩容，同时由后端统一控制业务状态。')

    add_heading(doc, '5.1.2 前端技术框架', 3)
    add_body(doc, '前端基于Vue 3、TypeScript和Vite构建，使用Vue Router管理页面路由，Pinia管理用户状态，Axios封装API请求。Element Plus提供表单、表格、弹窗和任务状态组件，AntV G6承担关系图谱渲染，ECharts用于总览和统计图。请求拦截器自动从localStorage读取JWT并写入Authorization头，响应拦截器统一处理CommonResult业务码。')

    add_heading(doc, '5.1.3 后端技术框架', 3)
    add_body(doc, '后端采用Java 17和Spring Boot 3.2.5。Spring Web提供REST接口，Spring Security和JJWT提供无状态JWT认证，MyBatis-Plus及JdbcTemplate访问PostgreSQL，Flyway维护数据库迁移，Spring AOP记录操作审计，springdoc-openapi生成接口文档。分析任务由固定大小线程池异步执行，长时间Worker调用设置超时并收集标准输出、标准错误和结构化错误。')

    add_heading(doc, '5.1.4 算法与数据基础设施', 3)
    add_body(doc, '结构化案例Worker和反欺诈Worker采用Python实现，二者使用独立业务规则和提示词，只复用通用的框架抽取和相似度算法。BGE-M3用于风险知识检索，bge-small-zh-v1.5用于中文向量召回，bge-reranker-v2-m3用于候选重排；最终案例排序由归一化近似图编辑距离确定。CoreNLP 4.5.10提供中文指代消解。')
    add_body(doc, 'PostgreSQL 16是业务数据和审计记录的权威存储，TuGraph 4.5.2是当前启用的图引擎。图写入使用Neo4j Java Driver兼容的Bolt协议访问BankGraph数据库；项目仍保留Neo4j和JanusGraph连接配置，用于兼容和扩展，但当前生产主路径由graph.active-engine=tugraph确定。')
    add_table(doc, '表5-1 当前项目主要技术组件',
              ['层次', '主要技术', '版本或实现', '主要职责'], [
                  ['前端', 'Vue、TypeScript、Vite', 'Vue 3.4.21；TypeScript 5.4.3；Vite 5.2.8', '页面交互、状态管理、API调用'],
                  ['可视化', 'AntV G6、ECharts', 'G6 5.1.1；ECharts 5.5.0', '关系图谱与统计图展示'],
                  ['后端', 'Java、Spring Boot', 'Java 17；Spring Boot 3.2.5', '业务流程、接口、安全和审计'],
                  ['数据访问', 'MyBatis-Plus、Flyway', '3.5.5；9.22.3', '数据访问、分页和结构迁移'],
                  ['业务数据库', 'PostgreSQL', '16', '案例、任务、知识结果和审计'],
                  ['图数据库', 'TuGraph', '4.5.2', '案例图存储、查询和关系遍历'],
                  ['算法Worker', 'Python', 'JSON-over-stdio', '报告、框架抽取和相似匹配'],
                  ['语义服务', 'CoreNLP与BGE模型', 'CoreNLP 4.5.10；BGE系列', '指代消解、向量召回和重排'],
              ], [1.0, 2.2, 3.2, 3.0])

    add_heading(doc, '5.1.5 部署与运行结构', 3)
    add_body(doc, '本地运行时，前端开发服务器、3030端口后端、PostgreSQL、TuGraph、CoreNLP和模型服务分别启动。生产环境可使用Nginx提供静态资源和/api反向代理，后端以JAR、Docker或systemd运行，Python Worker由后端按需拉起，不需要常驻HTTP端口。风险分析服务可独立监听18082端口。')
    add_body(doc, '基础设施可通过docker-compose-stack.yml统一部署：PostgreSQL使用5432端口，TuGraph Web、Bolt和RPC分别使用7070、7687和9090端口，CoreNLP映射到9002端口，BGE模型服务使用8080端口。数据库和图数据库使用外部持久化卷，容器停止不会删除业务数据。生产环境必须通过环境变量覆盖数据库密码、图数据库密码、JWT密钥和模型服务凭据。')

    add_heading(doc, '5.2 API总览', 2)
    add_heading(doc, '5.2.1 接口风格与基础约定', 3)
    add_body(doc, '业务接口采用REST风格，基础路径为/api/v1。GET用于查询，POST用于创建或执行动作，PUT用于完整更新或业务状态迁移，PATCH用于局部状态更新，DELETE用于删除。路径参数表示确定资源，例如/cases/{caseId}；查询参数用于分页和筛选；复杂输入使用JSON请求体；文件上传使用multipart/form-data。')
    add_body(doc, '接口默认使用UTF-8 JSON。时间字段采用ISO 8601字符串，金额由数据库高精度数值类型承载，案例、任务、图谱和分析运行均使用稳定业务标识。分页查询统一接受pageNum和pageSize，返回records、total、pageNum、pageSize和pages。')

    add_heading(doc, '5.2.2 身份认证与权限控制', 3)
    add_body(doc, '登录成功后返回access token和refresh token。调用受保护接口时，客户端在请求头中携带Authorization: Bearer <token>。后端使用无状态会话，不在服务器保存登录Session。JWT包含用户名、角色和银行编码，银行管理员只能访问绑定银行的数据，跨银行请求由CurrentUser范围检查拒绝。')
    add_body(doc, 'sadmin可访问系统管理和全部业务接口；badmin可访问案例上传、总览、案例、线索、研究分析和图谱接口，并受银行范围限制；事件元数据的读取对sadmin和badmin开放，系统元数据写操作和用户、角色、菜单管理仅允许sadmin。Swagger、OpenAPI、健康检查和登录接口无需JWT。')

    add_heading(doc, '5.2.3 统一响应与错误处理', 3)
    add_body(doc, '业务响应统一包装为CommonResult，字段包括code、message、data和timestamp。成功时code为200、message为success；业务错误可使用400、401、403、404、405、409或500。当前前端按响应体code判断业务成功，即使HTTP状态仍为200，只要code不是200就进入异常分支。调用方不能只依赖HTTP状态。')
    add_code(doc, '{\n  "code": 200,\n  "message": "success",\n  "data": { "jobId": "JOB-2026-001" },\n  "timestamp": "2026-08-29T10:30:00"\n}')

    add_heading(doc, '5.2.4 接口分组', 3)
    add_table(doc, '表5-2 API分组总览',
              ['接口分组', '基础路径', '主要资源', '主要用途'], [
                  ['认证接口', '/api/v1/auth', 'login、logout、me', '登录、退出和当前用户信息'],
                  ['分析任务', '/api/v1/analysis', '上传、jobs、worker-config', '案例上传、任务执行与Worker管理'],
                  ['案例处理', '/api/v1/case-processing', '阶段队列、actions、approve', '报告、抽取、相似和处理审批'],
                  ['案例管理', '/api/v1/cases', '案例、事件、信号、工作流', '案例查询、复核审批和知识成果'],
                  ['图谱接口', '/api/v1/graph', 'nodes、edges、subgraph、schema', '图查询、案例子图和图引擎状态'],
                  ['线索接口', '/api/v1/clues', 'clues、runs、analyze', '关联线索生成、维护和合并'],
                  ['研究分析', '/api/v1/research-analytics', 'datasets、runs、artifacts、knowledge', '任务2/3算法运行和资产管理'],
                  ['平台总览', '/api/v1/overview', 'overview', '系统统计和趋势'],
                  ['系统管理', '/api/v1/system', 'users、roles、menus、dict、event-metadata', '权限、菜单、字典和事件元数据'],
              ], [1.2, 2.5, 2.6, 3.5], font_size=10)

    add_heading(doc, '5.3 接口详细说明', 2)
    add_heading(doc, '5.3.1 认证与用户接口', 3)
    add_body(doc, '登录接口校验用户名和密码，返回JWT、角色、银行编码、权限和菜单树。前端据此生成导航并控制路由。当前用户接口用于恢复用户信息和更新昵称、联系方式等个人资料。')
    add_table(doc, '表5-3 认证与用户接口', ['方法', '路径', '输入', '输出及说明'], [
        ['POST', '/api/v1/auth/login', 'username、password', 'LoginResponse；返回token、refreshToken、roleCode、bankCode和menus'],
        ['POST', '/api/v1/auth/logout', '无', '退出成功；客户端同时清理本地令牌'],
        ['GET', '/api/v1/auth/me', 'Bearer JWT', '当前用户身份、角色、银行和菜单'],
        ['PUT', '/api/v1/auth/me', '可更新的用户资料字段', '更新后的当前用户信息'],
    ], [1.6, 2.8, 2.7, 3.9], font_size=10)

    add_heading(doc, '5.3.2 案例上传与分析任务接口', 3)
    add_body(doc, '案例上传接口返回uploadToken，任务创建接口再引用该凭据。单案例上传使用multipart/form-data；批处理上传使用一个CSV或XLSX文件。创建任务后调用start启动异步处理，通过任务列表和详情查看进度。')
    add_table(doc, '表5-4 案例上传与分析任务接口', ['方法', '路径', '关键输入', '功能'], [
        ['POST', '/api/v1/analysis/structured-case-files', 'basicInfo、customers、recognitionMode及可选文件', '上传单个反洗钱案例并返回uploadToken'],
        ['POST', '/api/v1/analysis/anti-fraud-case-files', 'basicInfo、customers、accounts、devices及事件链/文本', '上传单个反欺诈案例'],
        ['POST', '/api/v1/analysis/structured-case-batch-file', 'file、recognitionMode', '上传反洗钱CSV/XLSX批次'],
        ['POST', '/api/v1/analysis/anti-fraud-case-batch-file', 'file、recognitionMode', '上传反欺诈CSV/XLSX批次'],
        ['POST', '/api/v1/analysis/jobs', 'bankCode、workspaceId、jobType、scenarioCode、inputParams、steps', '创建分析任务'],
        ['GET', '/api/v1/analysis/jobs', 'status、jobType、scenarioCode、batchId、分页参数', '分页查询任务记录'],
        ['GET', '/api/v1/analysis/jobs/{jobId}', 'jobId', '查询任务、步骤、结果和错误'],
        ['PUT', '/api/v1/analysis/jobs/{jobId}/start', 'jobId', '启动待处理任务'],
        ['PUT', '/api/v1/analysis/jobs/{jobId}/cancel', 'jobId', '取消待处理或运行中任务'],
        ['PUT', '/api/v1/analysis/jobs/{jobId}/retry', 'jobId', '重试失败任务'],
        ['DELETE', '/api/v1/analysis/jobs/{jobId}', 'jobId', '删除任务及其独占处理成果'],
    ], [1.6, 3.3, 3.2, 2.9], font_size=9.5)

    add_heading(doc, '5.3.3 三项核心功能接口', 3)
    add_body(doc, '案例处理接口使用stage筛选业务队列，使用action触发三项核心功能。REPORT、FRAMEWORK和SIMILARITY只允许作用于相应前置状态，服务端逐案校验银行访问范围和案例状态。')
    add_table(doc, '表5-5 案例智能处理接口', ['方法', '路径', '关键输入', '功能'], [
        ['GET', '/api/v1/case-processing', 'stage、recognitionMode、scenarioCode、caseId、分页参数', '查询指定处理阶段的案例'],
        ['GET', '/api/v1/case-processing/{caseId}', 'caseId', '返回原始材料、报告、框架、图谱快照、相似结果和推荐等级'],
        ['POST', '/api/v1/case-processing/actions/{action}', 'action=REPORT/FRAMEWORK/SIMILARITY；caseIds', '批量执行可疑报告、框架抽取或相似匹配'],
        ['PUT', '/api/v1/case-processing/{caseId}/report', 'analysisText', '保存人工修改的报告并重新进入框架抽取'],
        ['PUT', '/api/v1/case-processing/{caseId}/approve', 'riskLevel、orderedSimilarCaseIds', '核定风险等级和相似案例排序'],
        ['GET', '/api/v1/case-processing/similarity-graph/approved', '无', '查询已审核案例之间的相似关系'],
    ], [1.6, 3.4, 3.2, 3.0], font_size=9.5)

    add_heading(doc, '5.3.4 案例管理与复核审批接口', 3)
    add_body(doc, '案例管理接口承担案例列表、详情、业务字段维护、状态流转和关联成果查询。案例通用复核状态主链与案例处理池状态分开：前者面向DRAFT、IN_REVIEW、PENDING_APPROVAL、APPROVED等业务状态，后者面向报告、抽取和相似匹配阶段。调用方应根据页面所属流程选择对应接口。')
    add_table(doc, '表5-6 案例管理主要接口', ['方法', '路径', '关键输入', '功能'], [
        ['GET', '/api/v1/cases', 'caseId、caseStatus、riskLevel、scenarioCode、分页参数', '查询案例列表'],
        ['GET', '/api/v1/cases/{caseId}', 'caseId', '查询案例基本信息和汇总指标'],
        ['PUT', '/api/v1/cases/{caseId}/overview', '业务领域、触发点、报告时间等', '更新案例概览字段'],
        ['PUT', '/api/v1/cases/{caseId}/submit', 'caseId', '提交案例进入复核'],
        ['PUT', '/api/v1/cases/{caseId}/review', 'reviewer、reviewOpinion、reviewResult', '记录复核结果'],
        ['PUT', '/api/v1/cases/{caseId}/approve', 'approver、approvalOpinion、approvalResult', '记录审批结果'],
        ['PUT', '/api/v1/cases/{caseId}/return', 'reason', '退回案例'],
        ['PUT', '/api/v1/cases/{caseId}/reassign', 'assignee', '重新分派案例'],
        ['GET', '/api/v1/cases/{caseId}/workflow', 'caseId', '查询状态流转和审计记录'],
        ['GET', '/api/v1/cases/{caseId}/events', 'caseId', '查询案例事件'],
        ['GET', '/api/v1/cases/{caseId}/signals', 'caseId', '查询案例关联风险信号'],
        ['GET', '/api/v1/cases/{caseId}/worker-result', 'caseId', '查询原始Worker结果'],
    ], [1.6, 3.4, 3.0, 3.3], font_size=9.5)

    add_heading(doc, '5.3.5 图谱查询与维护接口', 3)
    add_body(doc, '图谱接口以TuGraph为当前活动引擎，返回统一的nodes和edges结构。case-subgraph是案例结果页面和全景图谱的主要接口；通用query接口只允许受控查询语句，生产环境应限制写操作和高开销遍历。')
    add_table(doc, '表5-7 图谱接口', ['方法', '路径', '关键输入', '功能'], [
        ['GET', '/api/v1/graph/case-subgraph', 'caseId', '返回案例子图、图数据库信息和快照回退结果'],
        ['GET', '/api/v1/graph/subgraph', 'nodeId、depth', '查询指定节点的多跳子图'],
        ['GET', '/api/v1/graph/nodes', 'label、limit', '按标签查询节点'],
        ['GET', '/api/v1/graph/edges', 'type、limit', '按关系类型查询边'],
        ['POST', '/api/v1/graph/query', 'cypher', '执行受控Cypher查询'],
        ['GET', '/api/v1/graph/schema', '无', '查询节点和关系Schema'],
        ['POST', '/api/v1/graph/case-rebuild', 'caseId', '按关系型数据重建案例图'],
        ['POST', '/api/v1/graph/case-fallback', 'caseId', '创建案例最小回退子图'],
        ['GET', '/api/v1/graph/health', '无', '查询活动图引擎健康状态'],
        ['GET', '/api/v1/graph/engines', '无', '查询已配置图引擎状态'],
    ], [1.6, 3.4, 2.8, 3.6], font_size=9.5)

    add_heading(doc, '5.3.6 案例知识解释接口', 3)
    add_body(doc, '知识解释接口提供核心事件链、事项解释、证据、风险假设、技术发生和复核建议。刷新类接口重新投影当前案例知识成果，查询类接口返回当前有效快照，反馈接口保存人工结论。')
    add_table(doc, '表5-8 案例知识解释主要接口', ['方法', '路径', '功能'], [
        ['POST', '/api/v1/cases/{caseId}/matter-explanations/refresh', '刷新事项解释'],
        ['GET', '/api/v1/cases/{caseId}/matter-explanations', '查询当前事项解释'],
        ['GET', '/api/v1/cases/{caseId}/matter-explanations/{matterId}/evidence', '查询事项证据'],
        ['GET', '/api/v1/cases/{caseId}/core-chains/current', '查询当前核心事件链'],
        ['POST', '/api/v1/cases/{caseId}/core-chains/refresh', '刷新核心事件链'],
        ['GET', '/api/v1/cases/{caseId}/knowledge-explanation-chains', '按FULL或PANORAMA视图查询知识解释链'],
        ['GET', '/api/v1/cases/{caseId}/reasoning', '查询风险推理结果'],
        ['GET', '/api/v1/cases/{caseId}/technique-occurrences', '查询技术发生记录'],
        ['GET', '/api/v1/cases/{caseId}/review-suggestions', '查询复核建议'],
    ], [1.5, 5.1, 4.0], font_size=9.5)

    add_heading(doc, '5.3.7 线索与研究分析接口', 3)
    add_body(doc, '线索接口负责从图谱分析结果生成、维护和合并风险线索；研究分析接口负责运行任务2/3中的事件链、行为矩阵、风险扩散、增量学习、元路径、时序异常、超图和级联推理，并保存可复现的运行参数及分析资产。')
    add_table(doc, '表5-9 线索与研究分析主要接口', ['方法', '路径', '功能'], [
        ['POST', '/api/v1/clues/analyze', '执行图谱线索分析'],
        ['GET', '/api/v1/clues', '分页查询线索'],
        ['GET', '/api/v1/clues/runs', '查询线索分析运行记录'],
        ['PATCH', '/api/v1/clues/{clueId}/status', '更新线索状态'],
        ['POST', '/api/v1/clues/{clueId}/merge/{caseId}', '将线索合并到案例'],
        ['POST', '/api/v1/research-analytics/scopes/context', '构建跨案例分析上下文'],
        ['POST', '/api/v1/research-analytics/scopes/runs', '按范围执行研究分析'],
        ['POST', '/api/v1/research-analytics/cases/{caseId}/runs', '按单案例执行研究分析'],
        ['GET', '/api/v1/research-analytics/runs/{runId}', '查询运行详情'],
        ['GET', '/api/v1/research-analytics/runs/{runId}/artifacts', '查询分析资产'],
        ['PUT', '/api/v1/research-analytics/artifacts/{artifactId}/review', '审核分析资产'],
        ['GET', '/api/v1/research-analytics/knowledge', '查询审核后的研究知识'],
    ], [1.5, 5.3, 3.8], font_size=9.5)

    add_heading(doc, '5.3.8 平台总览与系统管理接口', 3)
    add_body(doc, '平台总览接口按时间和场景返回案例、任务和风险统计。系统管理接口提供用户、角色、菜单、字典和事件元数据维护，其中用户、角色、菜单和字典写操作仅限超级管理员。')
    add_table(doc, '表5-10 平台与系统管理接口', ['方法', '路径', '功能'], [
        ['GET', '/api/v1/overview', '按startDate、endDate和scenarioCode查询平台总览'],
        ['GET/POST/PUT/DELETE', '/api/v1/system/users', '用户查询、新增、更新和删除'],
        ['GET/POST/PUT/DELETE', '/api/v1/system/roles', '角色及权限配置'],
        ['GET/POST/PUT/DELETE', '/api/v1/system/menus', '菜单树配置'],
        ['GET/POST/PUT/DELETE', '/api/v1/system/dict/{dictType}', '业务字典管理'],
        ['GET/POST/PUT/DELETE', '/api/v1/system/event-metadata', 'AML/ANTI_FRAUD事件元数据管理'],
    ], [2.2, 4.8, 3.4], font_size=9.5)

    add_heading(doc, '5.4 系统内部组件接口', 2)
    add_heading(doc, '5.4.1 前端与后端接口', 3)
    add_body(doc, '前端所有请求通过统一Axios实例发送，自动附加Bearer JWT。后端返回CommonResult后，响应拦截器提取业务数据；code为401时清理本地令牌并跳转登录页，其他非200业务码转换为Promise异常。长时间案例处理接口可单独设置更长的客户端超时，任务型操作优先采用创建、启动和轮询模式。')

    add_heading(doc, '5.4.2 后端与案例Worker接口', 3)
    add_body(doc, 'Spring Boot通过ProcessBuilder启动framework_extraction.py或similarity_matching.py，在标准输入写入一个JSON请求，在标准输出读取最终JSON结果。核心请求字段包括jobId、sourcePath、processingMode、recognitionMode、targetStage、historyFiles、historySource和frameworkSettings。Worker可以输出进度事件，后端持久化结果后再向调用方返回任务状态。')
    add_code(doc, '{\n  "jobId": "JOB-2026-001",\n  "sourcePath": "/worker/uploads/<uploadToken>",\n  "processingMode": "SINGLE",\n  "recognitionMode": "NEW",\n  "targetStage": "FRAMEWORK",\n  "historySource": "POSTGRESQL",\n  "frameworkSettings": { "llmEnabled": true, "corenlpEnabled": true }\n}')

    add_heading(doc, '5.4.3 后端与数据存储接口', 3)
    add_body(doc, '业务数据通过MyBatis-Plus Mapper和JdbcTemplate写入PostgreSQL，数据库结构由根目录postgresql中的Flyway脚本唯一维护。案例、事件、处理池、案例库、相似排序、图谱快照、线索、研究运行和审计记录通过case_id、job_id、run_id等稳定标识关联。后端事务用于保证状态迁移与业务产物的一致性。')

    add_heading(doc, '5.4.4 后端与图数据库接口', 3)
    add_body(doc, '后端使用Bolt协议连接TuGraph的BankGraph数据库，通过参数受控的Cypher完成Schema检查、节点关系写入和子图查询。图写入前必须通过关系型投影合同校验；查询结果统一转换为id、labels、properties、source、target和type等前端可识别字段。图数据库凭据、地址和活动引擎均由环境变量配置。')

    add_heading(doc, '5.4.5 Worker与语义服务接口', 3)
    add_body(doc, 'Worker通过HTTP调用CoreNLP完成中文指代消解，通过兼容/v1/embeddings和/v1/rerank的模型接口调用BGE服务，通过OpenAI兼容的聊天模型接口生成报告和受约束抽取结果。模型服务地址、模型名称、设备、是否仅加载本地文件和超时参数由部署环境配置，页面请求不能覆盖生产模型凭据。')

    add_heading(doc, '5.5 异常与错误码说明', 2)
    add_heading(doc, '5.5.1 通用错误码', 3)
    add_table(doc, '表5-11 通用业务错误码', ['业务码', '含义', '典型场景', '调用方处理'], [
        ['200', '成功', '查询或操作完成', '读取data'],
        ['400', '请求错误', '缺少文件、字段非法、状态参数不支持', '修正请求后重试'],
        ['401', '未认证', 'JWT缺失、无效或过期', '重新登录'],
        ['403', '无权限', '角色不允许或跨银行访问', '停止操作并核对权限'],
        ['404', '资源不存在', '案例、用户、任务或图谱对象不存在', '刷新列表或核对标识'],
        ['405', '方法不支持', 'HTTP方法与接口不匹配', '按接口定义修改方法'],
        ['409', '状态或数据冲突', '阶段不匹配、唯一键冲突、非法状态迁移', '刷新状态并按当前阶段处理'],
        ['500', '服务器处理失败', '未预期异常或依赖服务失败', '记录时间和请求标识，联系运维'],
    ], [1.4, 1.8, 3.3, 3.5], font_size=10)

    add_heading(doc, '5.5.2 业务状态与重试原则', 3)
    add_body(doc, '三项核心功能均执行前置状态校验。收到409时，调用方应重新查询案例处理详情，不得盲目重复提交。报告、框架和相似匹配失败时案例保留在原阶段，可修复输入或依赖后重试；审批成功后重复审批应被拒绝。文件校验失败属于400，不创建可执行任务。')
    add_body(doc, '分析任务的PENDING、RUNNING、SUCCEEDED、FAILED和CANCELLED与案例的PENDING_REPORT等业务状态相互独立。任务成功表示Worker调用和持久化完成，案例下一步仍需在对应业务队列执行。客户端应分别展示任务状态和案例状态，避免把上传完成误认为案例审批完成。')

    add_heading(doc, '5.5.3 安全与运行异常', 3)
    add_body(doc, '上传接口必须限制文件类型、文件大小、文件数量和目标路径，服务端采用流式处理大文件，不读取整个MultipartFile到内存。Worker运行设置总时限，模型语义阶段设置独立时限；超时进程被终止并返回可识别错误。日志不得输出密码、JWT、模型密钥或完整敏感案例材料。')
    add_body(doc, '图数据库不可用时，案例查询可使用已保存图谱快照或最小回退图；模型不可用时，相似度算法可降级到结构化GED。降级必须在结果中明确标记，不得以静默方式冒充完整模型结果。')

    add_heading(doc, '5.6 API使用示例', 2)
    add_heading(doc, '5.6.1 登录并取得JWT', 3)
    add_code(doc, 'POST /api/v1/auth/login\nContent-Type: application/json\n\n{\n  "username": "badmin",\n  "password": "******"\n}')
    add_body(doc, '登录成功后，从data.token取得访问令牌。后续请求统一添加Authorization请求头。')
    add_code(doc, 'Authorization: Bearer <data.token>')

    add_heading(doc, '5.6.2 上传案例并创建任务', 3)
    add_body(doc, '第一步调用multipart上传接口取得uploadToken；第二步创建STRUCTURED任务，并在inputParams中写入工作流、处理方式、识别模式和上传凭据；第三步调用start接口启动。')
    add_code(doc, 'POST /api/v1/analysis/jobs\nContent-Type: application/json\nAuthorization: Bearer <token>\n\n{\n  "bankCode": "BANK_XA",\n  "workspaceId": 1,\n  "jobType": "STRUCTURED",\n  "jobName": "新增案例上传任务",\n  "scenarioCode": "AML",\n  "inputParams": "{\\"uploadToken\\":\\"<uploadToken>\\",\\"recognitionMode\\":\\"NEW\\"}",\n  "steps": [\n    {"stepName":"数据校验","stepType":"TEXT"},\n    {"stepName":"案例入库","stepType":"PERSIST"}\n  ]\n}')
    add_code(doc, 'PUT /api/v1/analysis/jobs/{jobId}/start\nAuthorization: Bearer <token>')

    add_heading(doc, '5.6.3 执行三项核心功能', 3)
    add_body(doc, '新增案例进入对应队列后，依次以REPORT、FRAMEWORK和SIMILARITY调用统一动作接口。每次调用成功后应重新查询案例详情，确认processingStage已经进入下一阶段。')
    add_code(doc, 'POST /api/v1/case-processing/actions/REPORT\nContent-Type: application/json\nAuthorization: Bearer <token>\n\n{ "caseIds": ["CASE-2026-001"] }')
    add_code(doc, 'POST /api/v1/case-processing/actions/FRAMEWORK\nContent-Type: application/json\nAuthorization: Bearer <token>\n\n{ "caseIds": ["CASE-2026-001"] }')
    add_code(doc, 'POST /api/v1/case-processing/actions/SIMILARITY\nContent-Type: application/json\nAuthorization: Bearer <token>\n\n{ "caseIds": ["CASE-2026-001"] }')

    add_heading(doc, '5.6.4 审批并查询案例图谱', 3)
    add_code(doc, 'PUT /api/v1/case-processing/CASE-2026-001/approve\nContent-Type: application/json\nAuthorization: Bearer <token>\n\n{\n  "riskLevel": "HIGH",\n  "orderedSimilarCaseIds": ["CASE-H-018", "CASE-H-102"]\n}')
    add_code(doc, 'GET /api/v1/graph/case-subgraph?caseId=CASE-2026-001\nAuthorization: Bearer <token>')
    add_body(doc, '图谱接口返回案例节点、实体节点、事件节点、证据节点和关系集合。前端根据节点类型和关系类型应用图形样式，并可继续请求知识解释链、关联线索分析或隐蔽风险挖掘任务。')

    add_heading(doc, '5.7 本章小结', 2)
    add_body(doc, '本章说明系统的前后端、算法Worker、关系数据库、图数据库与语义服务框架，并给出REST API分组、关键接口、内部协议、错误码和调用示例。接口采用身份与银行范围校验、统一响应、状态校验和审计记录，支撑案例上传、核心功能、复核审批和图谱分析。')


def main() -> None:
    if sha256(REFERENCE) != EXPECTED_REFERENCE_SHA256:
        raise SystemExit('Reference DOCX changed; fresh template distillation is required.')
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    working = OUTPUT.with_suffix('.working.docx')
    shutil.copy2(REFERENCE, working)
    doc = Document(working)
    original_text = [p.text for p in doc.paragraphs]
    configure_styles(doc)
    chapter_four(doc)
    chapter_five(doc)
    # Preserve reference A4 geometry and ensure all added body/table runs use
    # installed Chinese fonts.
    doc.save(OUTPUT)
    working.unlink(missing_ok=True)

    verify = Document(OUTPUT)
    if [p.text for p in verify.paragraphs[:len(original_text)]] != original_text:
        raise SystemExit('Reference text changed while appending chapters.')
    if len(verify.sections) != 1:
        raise SystemExit('Unexpected section count change.')
    print(OUTPUT)
    print(f'paragraphs={len(verify.paragraphs)} tables={len(verify.tables)} sections={len(verify.sections)}')


if __name__ == '__main__':
    main()
